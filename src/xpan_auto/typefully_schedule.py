"""Plan bounded Typefully schedule slots and record verified local GET results."""

from __future__ import annotations

from .deployment import verified_social_set_id

import argparse
import hashlib
import json
import os
import posixpath
import re
import stat
import tempfile
import xml.etree.ElementTree as ET
from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path
from typing import Any, Mapping
from zoneinfo import ZoneInfo
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from .copy_review import validate_review
from .publish_pipeline import FIXED_WORKBOOK_PATH, _headers, _row_identity, _row_values


SHANGHAI = ZoneInfo("Asia/Shanghai")
SLOTS = ("19:30", "19:45", "20:15", "20:30", "20:45", "21:00", "21:15", "21:30")
PROJECT_ROOT = Path(__file__).resolve().parents[2]


def _datetime(value: object, label: str) -> datetime:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{label} 缺少带时区的 ISO 时间")
    try:
        parsed = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
    except ValueError as exc:
        raise ValueError(f"{label} 不是有效 ISO 时间") from exc
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        raise ValueError(f"{label} 必须包含时区")
    return parsed


def _unwrap(value: Any, label: str) -> Any:
    if not isinstance(value, Mapping):
        return value
    if value.get("isError") or value.get("error") or value.get("success") is False:
        raise ValueError(f"{label} 快照包含 MCP 错误，不能作为完整快照")
    structured = value.get("structuredContent")
    if isinstance(structured, Mapping):
        return _unwrap(structured, label)
    return value


def _reject_incomplete_page(value: Any, label: str) -> None:
    if isinstance(value, Mapping):
        if value.get("truncated") is True:
            raise ValueError(f"{label} 快照已截断，不能作为完整成功快照")
        next_value = value.get("next")
        if next_value is not None and next_value != "" and not (isinstance(next_value, bool) and not next_value):
            raise ValueError(f"{label} 快照仍有下一页，不能作为完整成功快照")


def _draft_from_item(item: Any) -> Mapping[str, Any]:
    if not isinstance(item, Mapping):
        raise ValueError("MCP 快照中的 draft item 必须是对象")
    if item.get("kind") == "queue_slot" and "draft" in item and item["draft"] is None:
        return {}
    candidate = item.get("draft", item)
    if candidate is None:
        raise ValueError("非空快照项缺少 draft 对象")
    if not isinstance(candidate, Mapping):
        raise ValueError("MCP 快照中的 draft 必须是对象")
    if candidate.get("id", candidate.get("draft_id")) in (None, ""):
        raise ValueError("MCP 快照中的 draft 缺少有效 ID")
    return candidate


def _snapshot_drafts(snapshot: Any) -> list[Mapping[str, Any]]:
    snapshot = _unwrap(snapshot, "Typefully")
    _reject_incomplete_page(snapshot, "Typefully")
    if isinstance(snapshot, list):
        return [draft for item in snapshot if (draft := _draft_from_item(item))]
    if isinstance(snapshot, Mapping) and snapshot.get("id", snapshot.get("draft_id")) is not None:
        return [_draft_from_item(snapshot)]
    queue_part: Any = snapshot
    list_part: Any = None
    if isinstance(snapshot, Mapping) and "results" in snapshot:
        results_obj = _unwrap(snapshot, "list_drafts")
        _reject_incomplete_page(results_obj, "list_drafts")
        results = results_obj.get("results")
        if not isinstance(results, list):
            raise ValueError("list_drafts 快照必须包含 results 数组")
        return [_draft_from_item(item) for item in results]
    if isinstance(snapshot, Mapping) and ("queue" in snapshot or "list" in snapshot):
        if "queue" not in snapshot or "list" not in snapshot:
            raise ValueError("完整快照必须同时包含 queue 和 list keys")
        queue_part, list_part = snapshot["queue"], snapshot["list"]
    queue_part = _unwrap(queue_part, "queue")
    _reject_incomplete_page(queue_part, "queue")
    if not isinstance(queue_part, Mapping) or not isinstance(queue_part.get("days"), list):
        raise ValueError("queue 快照必须包含完整 days 数组")
    drafts: list[Mapping[str, Any]] = []
    for day in queue_part["days"]:
        if not isinstance(day, Mapping) or not isinstance(day.get("items"), list):
            raise ValueError("queue.days 每项必须包含 items 数组")
        drafts.extend(draft for item in day["items"] if (draft := _draft_from_item(item)))
    if list_part is not None:
        list_part = _unwrap(list_part, "list_drafts")
        _reject_incomplete_page(list_part, "list_drafts")
        if isinstance(list_part, list):
            results = list_part
        elif isinstance(list_part, Mapping) and isinstance(list_part.get("results"), list):
            results = list_part["results"]
        else:
            raise ValueError("list_drafts 快照必须是裸数组或包含 results 数组")
        drafts.extend(_draft_from_item(item) for item in results)
    return drafts


def _snapshot_identity(draft: Mapping[str, Any]) -> tuple[str, str, datetime | None]:
    draft_id = str(draft.get("id", draft.get("draft_id", ""))).strip()
    if not draft_id:
        raise ValueError("draft 快照缺少 ID")
    status = str(draft.get("status", "")).strip().lower()
    if status not in {"draft", "scheduled", "planned", "published", "publishing", "error"}:
        raise ValueError(f"draft {draft_id} 缺少有效状态")
    if status not in {"scheduled", "planned", "published"}:
        return draft_id, status, None
    raw_time = draft.get("scheduled_date")
    if status == "published" and not raw_time:
        raw_time = draft.get("published_at")
    if not raw_time:
        raise ValueError(f"{status} draft {draft_id} 缺少占位时间")
    return draft_id, status, _datetime(raw_time, f"draft {draft_id} 时间").astimezone(SHANGHAI)


def build_plan(target_date: str, now: str, existing_json: Any) -> dict[str, Any]:
    try:
        day = date.fromisoformat(target_date)
    except ValueError as exc:
        raise ValueError("--date 必须是 YYYY-MM-DD") from exc
    current = _datetime(now, "--now").astimezone(SHANGHAI)
    if not isinstance(existing_json, (Mapping, list)):
        raise ValueError("existing-json 必须是完整 queue/list 快照对象")
    drafts = _snapshot_drafts(existing_json)
    by_id: dict[str, tuple[str, datetime | None]] = {}
    for draft in drafts:
        draft_id, status, when = _snapshot_identity(draft)
        identity = (status, when)
        previous = by_id.get(draft_id)
        if previous is not None and previous != identity:
            raise ValueError(f"同一 draft ID {draft_id} 存在冲突状态或时间快照")
        by_id[draft_id] = identity

    occupied: dict[str, dict[str, str]] = {}
    for draft_id, (status, when) in by_id.items():
        if status not in {"scheduled", "planned", "published"} or when is None or when.date() != day:
            continue
        key = when.strftime("%H:%M")
        if key not in SLOTS:
            continue
        if key in occupied and occupied[key]["draft_id"] != draft_id:
            raise ValueError(f"固定槽 {key} 已有多个不同 draft 占用")
        occupied[key] = {"draft_id": draft_id, "status": f"occupied_{status}"}

    slots = []
    for clock in SLOTS:
        slot_dt = datetime.combine(day, time.fromisoformat(clock), tzinfo=SHANGHAI)
        if clock in occupied:
            state = occupied[clock]
            slots.append({"time": clock, "datetime": slot_dt.isoformat(), "status": state["status"], "draft_id": state["draft_id"]})
        elif slot_dt <= current:
            slots.append({"time": clock, "datetime": slot_dt.isoformat(), "status": "expired"})
        else:
            slots.append({"time": clock, "datetime": slot_dt.isoformat(), "status": "available"})
    return {"date": day.isoformat(), "timezone": "Asia/Shanghai", "daily_limit": 8, "slots": slots}


def _unwrap_draft(value: Any) -> Mapping[str, Any]:
    value = _unwrap(value, "draft GET")
    if isinstance(value, Mapping) and isinstance(value.get("draft"), Mapping):
        value = value["draft"]
    if not isinstance(value, Mapping):
        raise ValueError("--draft-json 必须包含完整 draft GET 对象")
    return value


def _draft_text_and_media(draft: Mapping[str, Any]) -> tuple[str, list[str], set[str]]:
    platforms = draft.get("platforms")
    if not isinstance(platforms, Mapping):
        raise ValueError("回读草稿缺少平台设置")
    x_data = platforms.get("x")
    if not isinstance(x_data, Mapping) or x_data.get("enabled") is not True:
        raise ValueError("回读草稿的 X 平台未开启")
    posts = x_data.get("posts")
    if not isinstance(posts, list) or len(posts) != 1 or not isinstance(posts[0], Mapping):
        raise ValueError("回读草稿必须只有一条 X post")
    post = posts[0]
    text = post.get("text")
    media_ids = post.get("media_ids")
    if not isinstance(text, str) or not isinstance(media_ids, list) or not all(isinstance(x, (str, int)) for x in media_ids):
        raise ValueError("回读草稿缺少正文或媒体 ID")
    enabled: set[str] = set()
    for key, value in platforms.items():
        if value is None:
            continue
        if not isinstance(value, Mapping):
            raise ValueError(f"回读草稿的平台 {key} 结构无效")
        if value.get("enabled") is True:
            enabled.add(str(key))
    if enabled != {"x"}:
        raise ValueError("回读草稿必须只有 X 平台开启")
    return text, [str(item) for item in media_ids], enabled


def _required_headers(headers: Mapping[str, int], names: tuple[str, ...]) -> None:
    missing = [name for name in names if name not in headers]
    if missing:
        raise ValueError("固定主表缺少必要列: " + ", ".join(missing))


def _cell_values(sheet, headers: Mapping[str, int], row: int) -> dict[str, Any]:
    return {name: sheet.cell(row, column).value for name, column in headers.items()}


def _style_signature(cell: Any) -> tuple[int, ...]:
    return tuple(cell._style) if cell.has_style else (0,) * 9


def record_schedule(*, row_index: int, row_number: int, draft_json: Any, expected_time: str,
                    social_set_id: int) -> dict[str, Any]:
    """Validate an externally read-back draft, then atomically record local schedule metadata."""
    if row_index < 1 or row_number < 2:
        raise ValueError("row-index 必须为正序号，row-number 必须是工作表真实行号")
    if social_set_id != verified_social_set_id():
        raise ValueError("social-set-id 必须与 XPAN_TYPEFULLY_SOCIAL_SET_ID 中已核对的账号一致")
    expected = _datetime(expected_time, "--expected-time").astimezone(SHANGHAI)
    if expected.strftime("%H:%M") not in SLOTS or expected.second or expected.microsecond:
        raise ValueError("expected-time 必须落在本地固定排期槽")
    draft = _unwrap_draft(draft_json)
    draft_id = str(draft.get("id", draft.get("draft_id", ""))).strip()
    if not draft_id:
        raise ValueError("回读草稿缺少 ID")
    if str(draft.get("social_set_id", draft.get("account_id", ""))) != str(social_set_id):
        raise ValueError("回读草稿 social/account ID 与已核对账号不符")
    if str(draft.get("status", "")).lower() != "scheduled":
        raise ValueError("只有 status=scheduled 的回读草稿可记账")
    actual = _datetime(draft.get("scheduled_date"), "回读草稿排期").astimezone(SHANGHAI)
    if actual != expected:
        raise ValueError("回读草稿 scheduled_date 与 expected-time 不一致")
    published_evidence = ("posted_at", "published_at", "x_post_published_at", "x_published_url")
    if any(draft.get(field) not in (None, "") for field in published_evidence):
        raise ValueError("草稿包含实际发布证据，不能记为 scheduled")
    updated_at = draft.get("updated_at")
    _datetime(updated_at, "回读草稿 updated_at")
    draft_text, media_ids, _ = _draft_text_and_media(draft)

    asset_dir = PROJECT_ROOT / "output" / "assets" / f"{row_index:04d}"
    workbook_path = FIXED_WORKBOOK_PATH
    if not workbook_path.is_file():
        raise ValueError("固定主表不存在")
    x_path = asset_dir / "x.txt"
    if not x_path.is_file():
        raise ValueError("目标行缺少已审阅的 x.txt")

    before_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    workbook = load_workbook(workbook_path)
    temporary: Path | None = None
    try:
        if "处理结果" not in workbook.sheetnames:
            raise ValueError("固定主表缺少工作表 处理结果")
        sheet = workbook["处理结果"]
        if not 2 <= row_number <= sheet.max_row:
            raise ValueError("真实 Excel 行号超出固定主表")
        headers = _headers(sheet)
        _required_headers(headers, (
            "序号", "资料标题", "文字介绍", "新夸克链接", "Typefully草稿ID", "Typefully状态",
            "Typefully更新时间", "Typefully错误信息", "X发布状态", "X排期时间", "X发布时间",
            "X帖子链接", "X错误信息",
        ))
        row = _cell_values(sheet, headers, row_number)
        serial, title, description, link = _row_identity(row, row_number)
        if serial != row_index:
            raise ValueError("序号与真实 Excel 行号不匹配")
        if str(row.get("Typefully草稿ID") or "") != draft_id:
            raise ValueError("回读 draft ID 与该行已记录草稿 ID 不匹配")
        if any(str(row.get(name) or "").strip() for name in ("X发布时间", "X帖子链接")):
            raise ValueError("该行已有发布结果，不能记为排期成功")
        if str(row.get("X发布状态") or "") == "已发布" or str(row.get("Typefully状态") or "").upper() == "PUBLISHED":
            raise ValueError("已发布行不能走排期记账")
        previous_slot = row.get("X排期时间")
        idempotent = (str(row.get("Typefully状态") or "").upper() == "SCHEDULED"
                and str(row.get("X发布状态") or "") == "已排期"
                and str(previous_slot or "") == expected.isoformat())
        if previous_slot and str(previous_slot) != expected.isoformat():
            raise ValueError("该行已有不同排期时间，拒绝覆盖")

        copy = x_path.read_text(encoding="utf-8")
        if copy != draft_text:
            raise ValueError("Typefully 回读正文与 x.txt 不完全一致")
        if copy.count(link) != 1 or not copy.rstrip().endswith(link) or copy.rstrip().splitlines()[-1].strip() != link:
            raise ValueError("x.txt 必须只包含一次本行夸克链接，且链接位于最后一行")
        validate_review(asset_dir, serial=serial, title=title, description=description, link=link)

        ledger_path = asset_dir / "typefully-upload-state.json"
        if not ledger_path.is_file():
            raise ValueError("缺少已保存 Typefully 上传状态，不能猜测媒体 ID")
        ledger = json.loads(ledger_path.read_text(encoding="utf-8"))
        ledger_media_id = str(ledger.get("media_id") or "")
        if (ledger.get("serial") != serial or ledger.get("excel_row_number") != row_number
                or str(ledger.get("social_set_id")) != str(social_set_id) or not ledger_media_id
                or media_ids != [ledger_media_id]):
            raise ValueError("回读媒体 ID 与该行已保存 Typefully 上传状态不匹配")
        media_file = str(ledger.get("media_file") or "")
        expected_hash = str(ledger.get("media_sha256") or "")
        media_path = (asset_dir / media_file).resolve()
        if (not media_file or not re.fullmatch(r"[0-9a-f]{64}", expected_hash)
                or not media_path.is_relative_to(asset_dir.resolve()) or not media_path.is_file()):
            raise ValueError("Typefully 上传状态缺少有效本地媒体或 SHA-256")
        with media_path.open("rb") as media_handle:
            actual_hash = hashlib.file_digest(media_handle, "sha256").hexdigest()
        if actual_hash != expected_hash:
            raise ValueError("本地媒体 SHA-256 与已核验的 Typefully 上传状态不一致")
        if idempotent:
            workbook.close()
            return {"serial": serial, "row_number": row_number, "draft_id": draft_id,
                    "status": "SCHEDULED", "idempotent": True}

        workbook_values = {
            "Typefully状态": "SCHEDULED",
            "Typefully更新时间": str(updated_at),
            "X发布状态": "已排期",
            "X排期时间": expected.isoformat(),
            "Typefully错误信息": "",
            "X错误信息": "",
        }
        workbook.close()
        temporary = _write_package_candidate(workbook_path, row_number, headers, workbook_values)
        source_mode = stat.S_IMODE(workbook_path.stat().st_mode)
        os.chmod(temporary, source_mode)
        if stat.S_IMODE(temporary.stat().st_mode) != source_mode:
            raise ValueError("候选主表未保留源文件权限 mode")
        _verify_candidate(temporary, workbook_path, row_number, headers, workbook_values)
        if hashlib.sha256(workbook_path.read_bytes()).hexdigest() != before_hash:
            raise ValueError("固定主表在检查期间发生并发变化，取消原子替换")
        os.replace(temporary, workbook_path)
        temporary = None
        return {"serial": serial, "row_number": row_number, "draft_id": draft_id,
                "status": "SCHEDULED", "idempotent": False}
    finally:
        try:
            workbook.close()
        except Exception:
            pass
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def _worksheet_part(archive: ZipFile, title: str) -> str:
    workbook_xml = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships_xml = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    relationship_targets = {
        rel.attrib["Id"]: rel.attrib["Target"]
        for rel in relationships_xml.findall(f"{{{package_ns}}}Relationship")
    }
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    sheets = workbook_xml.find(f"{{{main_ns}}}sheets")
    if sheets is None:
        raise ValueError("主表 XLSX 缺少 worksheets 列表")
    for sheet in sheets.findall(f"{{{main_ns}}}sheet"):
        if sheet.attrib.get("name") == title:
            relationship_id = sheet.attrib.get(f"{{{rel_ns}}}id")
            target = relationship_targets.get(relationship_id or "")
            if not target:
                break
            return posixpath.normpath(target.lstrip("/") if target.startswith("/")
                                      else posixpath.join("xl", target))
    raise ValueError(f"无法定位 XLSX 工作表 {title}")


def _register_xml_namespaces(payload: bytes) -> None:
    for _, namespace in ET.iterparse(BytesIO(payload), events=("start-ns",)):
        prefix, uri = namespace
        ET.register_namespace(prefix or "", uri)


def _patch_sheet_xml(payload: bytes, row_number: int, headers: Mapping[str, int],
                     values: Mapping[str, str]) -> bytes:
    _register_xml_namespaces(payload)
    root = ET.fromstring(payload)
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    xml_ns = "http://www.w3.org/XML/1998/namespace"
    sheet_data = root.find(f"{{{main_ns}}}sheetData")
    if sheet_data is None:
        raise ValueError("处理结果 worksheet 缺少 sheetData")
    row = next((element for element in sheet_data.findall(f"{{{main_ns}}}row")
                if element.attrib.get("r") == str(row_number)), None)
    if row is None:
        raise ValueError("处理结果 worksheet XML 缺少目标行")

    def column_number(cell: ET.Element) -> int:
        match = re.match(r"([A-Z]+)", cell.attrib.get("r", ""))
        return column_index_from_string(match.group(1)) if match else 0

    for name, value in values.items():
        column = headers[name]
        coordinate = f"{get_column_letter(column)}{row_number}"
        cell = next((element for element in row.findall(f"{{{main_ns}}}c")
                     if element.attrib.get("r") == coordinate), None)
        if cell is None:
            cell = ET.Element(f"{{{main_ns}}}c", {"r": coordinate})
            cells = list(row.findall(f"{{{main_ns}}}c"))
            insert_at = next((list(row).index(other) for other in cells if column_number(other) > column), len(row))
            row.insert(insert_at, cell)
        for child in list(cell):
            cell.remove(child)
        if value in (None, ""):
            cell.attrib.pop("t", None)
            continue
        cell.attrib["t"] = "inlineStr"
        inline = ET.SubElement(cell, f"{{{main_ns}}}is")
        text = ET.SubElement(inline, f"{{{main_ns}}}t")
        text.text = str(value)
        if text.text[:1].isspace() or text.text[-1:].isspace():
            text.set(f"{{{xml_ns}}}space", "preserve")
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _write_package_candidate(original: Path, row_number: int, headers: Mapping[str, int],
                             values: Mapping[str, str]) -> Path:
    temporary: Path | None = None
    try:
        with ZipFile(original) as source:
            worksheet = _worksheet_part(source, "处理结果")
            patched = _patch_sheet_xml(source.read(worksheet), row_number, headers, values)
            with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
                temporary = Path(handle.name)
            with ZipFile(temporary, "w") as target:
                target.comment = source.comment
                for info in source.infolist():
                    contents = patched if info.filename == worksheet else source.read(info.filename)
                    target.writestr(info, contents)
        return temporary
    except Exception:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        raise


def _verify_candidate(candidate: Path, original: Path, target_row: int,
                      headers: Mapping[str, int], allowed_values: Mapping[str, str]) -> None:
    with ZipFile(original) as before_zip, ZipFile(candidate) as candidate_zip:
        if candidate_zip.testzip() is not None:
            raise ValueError("候选主表 ZIP 完整性校验失败")
        original_names = before_zip.namelist()
        if original_names != candidate_zip.namelist():
            raise ValueError("候选主表改变了 ZIP 部件集合或顺序")
        worksheet = _worksheet_part(before_zip, "处理结果")
        for name in original_names:
            if name != worksheet and before_zip.read(name) != candidate_zip.read(name):
                raise ValueError(f"候选主表改动了受保护 ZIP 部件: {name}")

    old = load_workbook(original)
    new = load_workbook(candidate)
    allowed = {headers[name]: value for name, value in allowed_values.items()}
    try:
        if old.sheetnames != new.sheetnames or "处理结果" not in old.sheetnames:
            raise ValueError("保存候选改变工作表集合或缺少处理结果")
        for title in old.sheetnames:
            old_sheet, new_sheet = old[title], new[title]
            if (old_sheet.max_row, old_sheet.max_column) != (new_sheet.max_row, new_sheet.max_column):
                raise ValueError(f"保存候选改变了工作表范围: {title}")
            for row in range(1, old_sheet.max_row + 1):
                for column in range(1, old_sheet.max_column + 1):
                    before, after = old_sheet.cell(row, column), new_sheet.cell(row, column)
                    if _style_signature(before) != _style_signature(after):
                        raise ValueError(f"保存候选改变了单元格样式: sheet={title}, row={row}, column={column}")
                    permitted = title == "处理结果" and row == target_row and column in allowed
                    if permitted:
                        expected = allowed[column]
                        if (after.value or "") != expected:
                            raise ValueError("候选主表目标字段与预期值不一致")
                    elif before.value != after.value:
                        raise ValueError("保存候选修改了白名单之外的工作簿数据")
    finally:
        old.close()
        new.close()


def _read_json(path: Path) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError(f"无法读取有效 JSON: {path}") from exc


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="仅计算本地 Typefully 排期计划，或记录经真实 MCP GET 回读核验的排期；本工具不调用 API、不发送、不操作浏览器。"
    )
    commands = parser.add_subparsers(dest="command", required=True)
    plan = commands.add_parser("plan", help="从完整 queue/list 成功快照列出固定槽状态")
    plan.add_argument("--date", required=True, help="目标日期 YYYY-MM-DD；可预览未来日期，日常排期传当天日期")
    plan.add_argument("--now", help="带时区 ISO 当前时间；默认当前上海时间")
    plan.add_argument("--existing-json", required=True, type=Path,
                      help="本次 MCP get_queue 与 list_drafts(status=scheduled,planned,published) 的完整 JSON 成功快照；接受 {queue:{days:[{items:[{draft:...}]}]},list:[draft,...]}、list.results、单独 queue days、裸数组或单条 draft。queue_slot 的 draft:null 视为空槽；拒绝坏对象、缺少完整结果、next 非空或 truncated=true 的分页快照。")
    record = commands.add_parser("record", help="记账真实 MCP GET 回读后的排期成功对象")
    record.add_argument("--row-index", required=True, type=int, help="资料序号")
    record.add_argument("--row-number", required=True, type=int, help="真实 Excel 行号")
    record.add_argument("--draft-json", required=True, type=Path,
                        help="真实 Typefully MCP GET 同一草稿后保存的完整对象 JSON（支持 structuredContent 包装）；本地工具不读取 API")
    record.add_argument("--expected-time", required=True, help="通过实际排期 UI 设置并由 MCP GET 回读确认的带时区 ISO 时间")
    record.add_argument("--social-set-id", required=True, type=int, help="已核对的 X Social Set ID，须与 XPAN_TYPEFULLY_SOCIAL_SET_ID 一致")
    args = parser.parse_args(argv)
    try:
        if args.command == "plan":
            now = args.now or datetime.now(SHANGHAI).isoformat()
            print(json.dumps(build_plan(args.date, now, _read_json(args.existing_json)), ensure_ascii=False, indent=2))
        else:
            result = record_schedule(row_index=args.row_index, row_number=args.row_number,
                                     draft_json=_read_json(args.draft_json), expected_time=args.expected_time,
                                     social_set_id=args.social_set_id)
            print(json.dumps(result, ensure_ascii=False))
    except (ValueError, OSError, json.JSONDecodeError) as exc:
        parser.error(str(exc))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
