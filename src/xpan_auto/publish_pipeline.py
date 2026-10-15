from __future__ import annotations

import argparse
import json
import subprocess
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook

from .copy_review import CopyReviewError, validate_review, validate_review_batch
from .telegram_quality import inspect_telegram_outputs, write_quality_report
from .video_quality import inspect_publish_video
from .runtime import resolve_project_path


PROJECT_ROOT = Path(__file__).resolve().parents[2]
FIXED_WORKBOOK_PATH = (PROJECT_ROOT / "output" / "quark_links.typefully.xlsx").resolve()
TELEGRAM_SCRIPT_ROOT = PROJECT_ROOT / ".codex" / "skills" / "telegram-auto-publish" / "scripts"
TELEGRAM_SCRIPTS = {
    "extract": TELEGRAM_SCRIPT_ROOT / "extract_quark_project.py",
    "prepare": TELEGRAM_SCRIPT_ROOT / "prepare_telegram_asset.py",
    "caption": TELEGRAM_SCRIPT_ROOT / "build_telegram_caption.py",
    "send": TELEGRAM_SCRIPT_ROOT / "send_telegram.py",
}

TITLE_COLUMNS = ("资料标题", "标题", "名称")
DESCRIPTION_COLUMNS = ("文字介绍", "简介", "描述")
LINK_COLUMNS = ("新夸克链接", "新分享链接", "分享链接", "链接")
PROCESS_STATUS_COLUMNS = ("处理状态", "状态", "status")
TELEGRAM_STATUS = "Telegram发布状态"
TELEGRAM_MESSAGE_ID = "Telegram消息ID"
TELEGRAM_PUBLISHED_AT = "Telegram发布时间"
TELEGRAM_ERROR = "Telegram错误信息"
X_ASSET_STATUS = "X资料状态"
X_ASSET_READY = "资料已准备完成"
X_ASSET_REVIEW_PENDING = "文案待审阅"
X_PUBLISH_STATUS = "X发布状态"
X_ERROR = "X错误信息"
TYPEFULLY_DRAFT_ID = "Typefully草稿ID"
TYPEFULLY_STATUS = "Typefully状态"
TYPEFULLY_ERROR = "Typefully错误信息"
PIPELINE_COLUMNS = ("统一发布状态", "统一发布错误信息", "统一发布更新时间")
MEDIA_COLUMNS = ("媒体类型", "媒体路径", "视频", "原视频", "Telegram视频")

TELEGRAM_SENT_STATUSES = {"已发送", "SENT", "TELEGRAM_SENT"}
TELEGRAM_SENDING_STATUSES = {"发送中", "SENDING"}
FAILED_STATUSES = {"FAILED", "TELEGRAM_FAILED"}
X_DONE_STATUSES = {"已排期", "已发布"}
TYPEFULLY_DONE_STATUSES = {"DRAFT_CREATED", "SCHEDULED", "PUBLISHED"}
TELEGRAM_MAX_REGENERATIONS = 2
TELEGRAM_MEDIA_POLICY_VERSION = 6
PUBLISH_MEDIA_POLICY_VERSION = 10
TELEGRAM_BOT_API_MAX_MEDIA_BYTES = 50_000_000
TELEGRAM_PREVIEWABLE_EXTENSIONS = {
    ".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp",
    ".mp4", ".avi", ".mov", ".mkv", ".webm", ".flv", ".m4v",
    ".pdf", ".doc", ".docx", ".ppt", ".pptx", ".xls", ".xlsx",
    ".txt", ".md", ".csv", ".json", ".html", ".htm", ".rtf",
    ".xml", ".yaml", ".yml", ".log", ".ini", ".sql",
}


class PublishPipelineError(RuntimeError):
    """Raised when a channel preparation or publishing step cannot continue."""


class TelegramSendUncertainError(PublishPipelineError):
    """Raised when an upload started but its Telegram result is unknown."""


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _resolve_path(value: str | Path) -> Path:
    return resolve_project_path(value).resolve()


def _value(row: Mapping[str, object], names: Iterable[str]) -> str:
    for name in names:
        value = row.get(name)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def _first_nonempty_line(value: str) -> str:
    for line in value.splitlines():
        if line.strip():
            return line.strip()
    return ""


def _headers(sheet) -> dict[str, int]:
    return {
        str(cell.value): index
        for index, cell in enumerate(sheet[1], start=1)
        if cell.value is not None
    }


def _ensure_columns(sheet, names: Iterable[str]) -> dict[str, int]:
    headers = _headers(sheet)
    for name in names:
        if name not in headers:
            column = sheet.max_column + 1
            sheet.cell(row=1, column=column, value=name)
            headers[name] = column
    return headers


def _row_values(sheet, headers: Mapping[str, int], row_number: int) -> dict[str, object]:
    return {name: sheet.cell(row=row_number, column=column).value for name, column in headers.items()}


def _load_candidates(
    workbook_path: Path,
    *,
    only_status: set[str],
    limit: int,
    run_all: bool,
    force_reprocess: bool = False,
) -> list[int]:
    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    headers = _headers(sheet)
    status_column = next((headers[name] for name in PROCESS_STATUS_COLUMNS if name in headers), None)
    if status_column is None:
        raise PublishPipelineError("Excel 中缺少处理状态列")

    candidates: list[int] = []
    for row_number in range(2, sheet.max_row + 1):
        status = str(sheet.cell(row=row_number, column=status_column).value or "")
        if status not in only_status:
            continue
        telegram_status = str(sheet.cell(row=row_number, column=headers[TELEGRAM_STATUS]).value or "") if TELEGRAM_STATUS in headers else ""
        telegram_message_id = str(sheet.cell(row=row_number, column=headers[TELEGRAM_MESSAGE_ID]).value or "") if TELEGRAM_MESSAGE_ID in headers else ""
        x_publish_status = str(sheet.cell(row=row_number, column=headers[X_PUBLISH_STATUS]).value or "") if X_PUBLISH_STATUS in headers else ""
        typefully_status = str(sheet.cell(row=row_number, column=headers[TYPEFULLY_STATUS]).value or "") if TYPEFULLY_STATUS in headers else ""
        telegram_done = telegram_status in TELEGRAM_SENT_STATUSES or bool(telegram_message_id)
        telegram_in_flight = telegram_status in TELEGRAM_SENDING_STATUSES
        x_done = (
            x_publish_status in X_DONE_STATUSES
            or typefully_status in TYPEFULLY_DONE_STATUSES
        )
        if not force_reprocess and (telegram_done or telegram_in_flight) and x_done:
            continue
        candidates.append(row_number)
    workbook.close()
    return candidates if run_all else candidates[:limit]


def _update_cells(workbook_path: Path, row_number: int, values: Mapping[str, object]) -> None:
    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    headers = _ensure_columns(sheet, values.keys())
    for name, value in values.items():
        sheet.cell(row=row_number, column=headers[name], value=value)
    workbook.save(workbook_path)


def _write_image_path(workbook_path: Path, row_number: int, serial: int, image: Path) -> None:
    """Update the verified row by header without Windows/WPS automation."""
    if not _valid_file(image):
        raise PublishPipelineError(f"图片文件不存在或为空：{image}")
    workbook = load_workbook(workbook_path)
    temporary = None
    try:
        sheet = workbook["处理结果"]
        headers = _headers(sheet)
        if "序号" not in headers or "图片" not in headers:
            raise PublishPipelineError("工作表缺少序号或图片列")
        if str(sheet.cell(row_number, headers["序号"]).value) != str(serial):
            raise PublishPipelineError("图片回写行号与序号不匹配")
        sheet.cell(row_number, headers["图片"], str(image.resolve()))
        with tempfile.NamedTemporaryFile(suffix=".xlsx", dir=workbook_path.parent, delete=False) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        temporary.replace(workbook_path)
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def _run_command(command: Sequence[str]) -> str:
    try:
        result = subprocess.run(
            list(command),
            cwd=PROJECT_ROOT,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            check=False,
        )
    except OSError as exc:
        raise PublishPipelineError(f"无法运行命令 {command[0]}: {exc}") from exc
    if result.returncode != 0:
        details = (result.stderr or result.stdout or "").strip()
        raise PublishPipelineError(f"命令失败（退出码 {result.returncode}）：{details[-1200:]}")
    return result.stdout


def _last_json(stdout: str) -> dict[str, Any]:
    for line in reversed(stdout.splitlines()):
        line = line.strip()
        if not line.startswith("{"):
            continue
        try:
            payload = json.loads(line)
        except json.JSONDecodeError:
            continue
        if isinstance(payload, dict):
            return payload
    return {}


def _valid_file(path: Path) -> bool:
    return path.is_file() and path.stat().st_size > 0


def _details_ready(
    path: Path,
    *,
    row_index: int | None = None,
    title: str = "",
    share_link: str = "",
) -> bool:
    """Accept only details produced by the current real-media selection policy."""
    if not _valid_file(path):
        return False
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError):
        return False
    if (
        not isinstance(payload, dict)
        or payload.get("media_policy_version") != TELEGRAM_MEDIA_POLICY_VERSION
        or payload.get("publish_media_policy_version") != PUBLISH_MEDIA_POLICY_VERSION
    ):
        return False
    if row_index is not None:
        try:
            if int(payload.get("row_index")) != row_index:
                return False
        except (TypeError, ValueError):
            return False
    if title and str(payload.get("title") or "").strip() != title.strip():
        return False
    if share_link and _first_nonempty_line(str(payload.get("source_share_link") or "")) != share_link:
        return False
    if not isinstance(payload.get("items"), list) or not payload["items"]:
        return False
    summary = payload.get("summary") if isinstance(payload.get("summary"), dict) else {}
    if summary.get("truncated") is True:
        return False

    media = payload.get("media")
    if not isinstance(media, dict):
        return False
    kind = str(media.get("kind") or "")
    if kind == "project-video":
        value = media.get("publish_path") or media.get("original_path")
        if not value:
            return False
        candidate = resolve_project_path(str(value), base=path.parent)
        if not _valid_file(candidate):
            return False
    if kind == "no-project-media":
        summary = payload.get("summary") if isinstance(payload.get("summary"), dict) else {}
        extension_counts = summary.get("extension_counts") if isinstance(summary.get("extension_counts"), dict) else {}
        for extension, count in extension_counts.items():
            try:
                available = int(count or 0) > 0
            except (TypeError, ValueError):
                return False
            if available and str(extension).lower() in TELEGRAM_PREVIEWABLE_EXTENSIONS:
                return False
        for item in payload.get("items", ()) or ():
            if not isinstance(item, dict) or str(item.get("type") or "") != "file":
                continue
            extension = str(item.get("extension") or Path(str(item.get("name") or "")).suffix).lower()
            if extension in TELEGRAM_PREVIEWABLE_EXTENSIONS:
                return False
        return True
    if kind not in {
        "project-video",
        "existing-local",
        "project-image",
        "video-preview",
        "video-frame",
        "video-contact-sheet",
        "document-preview",
    }:
        return False
    source_value = media.get("source_path")
    if not source_value:
        return False
    source_path = resolve_project_path(str(source_value), base=path.parent)
    return _valid_file(source_path)


def _row_identity(row: Mapping[str, object], row_number: int) -> tuple[int, str, str, str]:
    raw_index = row.get("序号")
    try:
        row_index = int(raw_index) if raw_index is not None else row_number - 1
    except (TypeError, ValueError):
        row_index = row_number - 1
    title = _value(row, TITLE_COLUMNS)
    description = _value(row, DESCRIPTION_COLUMNS)
    link = _first_nonempty_line(_value(row, LINK_COLUMNS))
    if row_index < 1:
        raise PublishPipelineError(f"Excel 行 {row_number} 的资料序号无效：{row_index}")
    if not title:
        raise PublishPipelineError(f"Excel 行 {row_number} 缺少资料标题")
    if not link:
        raise PublishPipelineError(f"Excel 行 {row_number} 缺少新夸克链接")
    return row_index, title, description, link


def _prepare_telegram_row(
    workbook_path: Path,
    row_number: int,
    row: Mapping[str, object],
    *,
    image_root: Path,
    quark_db_path: Path,
    cookie_file: Path,
    send: bool,
    force_send: bool,
    image_column: int,
    quark_auto_save_dir: Path = PROJECT_ROOT / "work" / "quark-auto-save",
) -> dict[str, Any]:
    row_index, title, description, link = _row_identity(row, row_number)
    asset_dir = image_root / f"{row_index:04d}"
    asset_dir.mkdir(parents=True, exist_ok=True)
    details_file = asset_dir / "telegram-project-details.json"
    caption_file = asset_dir / "telegram-caption.html"
    x_copy_file = asset_dir / "x.txt"

    details_ready = _details_ready(
        details_file,
        row_index=row_index,
        title=title,
        share_link=link,
    )
    if not details_ready:
        _run_command(
            [
                sys.executable,
                str(TELEGRAM_SCRIPTS["extract"]),
                "--row-index",
                str(row_index),
                "--share-link",
                link,
                "--title",
                title,
                "--db",
                str(quark_db_path),
                "--cookie-file",
                str(cookie_file),
                "--asset-root",
                str(image_root),
                "--quark-auto-save-dir",
                str(quark_auto_save_dir),
            ]
        )

        if not _details_ready(
            details_file,
            row_index=row_index,
            title=title,
            share_link=link,
        ):
            try:
                failed_details = json.loads(details_file.read_text(encoding="utf-8"))
            except (OSError, UnicodeDecodeError, json.JSONDecodeError):
                failed_details = {}
            failed_media = failed_details.get("media") if isinstance(failed_details, dict) else None
            if isinstance(failed_media, dict) and failed_media.get("kind") == "media-preview-required":
                diagnostics = failed_media.get("diagnostics")
                reason = str(diagnostics[0]) if isinstance(diagnostics, list) and diagnostics else str(failed_media.get("error") or "")
                if "download file size limit" in reason:
                    items = failed_details.get("items") if isinstance(failed_details, dict) else None
                    largest = max((int(item.get("size_bytes") or 0) for item in items
                                   if isinstance(item, dict) and str(item.get("type")) == "file"), default=0) if isinstance(items, list) else 0
                    raise PublishPipelineError(
                        f"夸克拒绝下载源文档（文件大小限制，最大文件 {largest:,} bytes）；"
                        "缺少真实文档预览，不能生成封面或发布"
                    )
                raise PublishPipelineError(f"真实媒体预览不可用：{reason[:350]}")
            raise PublishPipelineError("项目媒体提取完成后仍未满足视频优先/图片兜底策略")
    try:
        details = json.loads(details_file.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise PublishPipelineError(f"项目详情无法读取：{details_file}") from exc
    media = details.get("media") if isinstance(details, dict) else None
    if not isinstance(media, dict):
        raise PublishPipelineError("项目详情缺少发布媒体")
    media_kind = str(media.get("kind") or "")

    # A retry must not replace human-reviewed copy with a fresh seed. Any
    # changed source/detail hash is caught by validate_review before sending.
    if not send and not (asset_dir / "copy-review.json").is_file():
        caption_command = [
            sys.executable,
            str(TELEGRAM_SCRIPTS["caption"]),
            "--title",
            title,
            "--description",
            description,
            "--link",
            link,
            "--details-file",
            str(details_file),
            "--output",
            str(caption_file),
            "--x-output",
            str(x_copy_file),
        ]
        _run_command(caption_command)
    if not _valid_file(caption_file):
        raise CopyReviewError(f"Telegram 文案不存在；先准备并审阅：{caption_file}")
    if not _valid_file(x_copy_file):
        raise CopyReviewError(f"X 文案不存在；先准备并审阅：{x_copy_file}")

    if media_kind == "project-video":
        original_path = resolve_project_path(
            str(media.get("publish_path") or media.get("original_path") or ""), base=asset_dir,
        )
        media_path = original_path
        if not _valid_file(original_path):
            raise PublishPipelineError(f"最终发布视频不存在：{original_path}")
        report = inspect_publish_video(
            media_path,
            max_duration_seconds=600.0,
            max_size_bytes=TELEGRAM_BOT_API_MAX_MEDIA_BYTES - 1,
        )
        quality_path = asset_dir / "video-quality.json"
        try:
            quality_payload = json.loads(quality_path.read_text(encoding="utf-8"))
        except (OSError, UnicodeDecodeError, json.JSONDecodeError):
            quality_payload = {}
        if not isinstance(quality_payload, dict):
            quality_payload = {}
        quality_payload.update(
            {
                "passed": report.passed,
                "telegram_route": "bot-api",
                "pipeline_final": report.__dict__,
            }
        )
        quality_path.write_text(
            json.dumps(quality_payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        if not report.passed:
            raise PublishPipelineError(
                f"最终发布视频质检未通过，已阻断发送：{'；'.join(report.issues)}"
            )
        media_size = media_path.stat().st_size
        if media_size >= TELEGRAM_BOT_API_MAX_MEDIA_BYTES:
            raise PublishPipelineError(
                "最终发布视频必须严格小于 50,000,000 bytes，不能转入 Web K 中转"
            )
        telegram_route = "bot-api"
        x_media = original_path
        media_type = "video"
    else:
        cover_file = asset_dir / "cover-v1-information-card.png"
        quality_file = asset_dir / "telegram-quality.json"
        if not _valid_file(cover_file):
            _run_command(
                [
                    sys.executable,
                    str(TELEGRAM_SCRIPTS["prepare"]),
                    "--row-index",
                    str(row_index),
                    "--asset-root",
                    str(image_root),
                    "--title",
                    title,
                    "--description",
                    description,
                    "--details-file",
                    str(details_file),
                ]
            )
        if not _valid_file(cover_file):
            raise PublishPipelineError(f"兜底封面生成后仍不存在：{cover_file}")

        quality_report = None
        for attempt in range(1, TELEGRAM_MAX_REGENERATIONS + 2):
            quality_report = inspect_telegram_outputs(
                image_path=cover_file,
                caption_path=caption_file,
                asset_metadata_path=asset_dir / "telegram-asset.json",
                details_path=details_file,
                title=title,
                link=link,
            )
            write_quality_report(quality_file, quality_report, attempt=attempt)
            if quality_report.passed:
                break
            if attempt > TELEGRAM_MAX_REGENERATIONS:
                problems = "；".join(quality_report.issues)
                raise PublishPipelineError(f"Telegram 兜底图片质检未通过，已阻断发送：{problems}")
            _run_command(
                [
                    sys.executable,
                    str(TELEGRAM_SCRIPTS["prepare"]),
                    "--row-index",
                    str(row_index),
                    "--asset-root",
                    str(image_root),
                    "--title",
                    title,
                    "--description",
                    description,
                    "--details-file",
                    str(details_file),
                    "--force-regenerate",
                ]
            )
        if quality_report is None or not quality_report.passed:
            raise PublishPipelineError("Telegram 兜底图片未生成通过报告，已阻断发送")
        _write_image_path(workbook_path, row_number, row_index, cover_file)
        media_path = cover_file
        original_path = cover_file
        x_media = cover_file
        media_size = cover_file.stat().st_size
        if media_size >= TELEGRAM_BOT_API_MAX_MEDIA_BYTES:
            raise PublishPipelineError(
                "最终图片必须严格小于 50,000,000 bytes；Telegram 只允许 Bot API，不使用 Web K 中转"
            )
        telegram_route = "bot-api"
        media_type = "image"

    if not _valid_file(x_media) or not _valid_file(x_copy_file):
        raise PublishPipelineError("X 发布资料不完整，必须同时存在媒体和 x.txt")
    x_asset_status = X_ASSET_REVIEW_PENDING
    if send:
        # A reviewed X draft is required even when Telegram was sent earlier.
        validate_review(asset_dir, serial=row_index, title=title, description=description, link=link)
        x_asset_status = X_ASSET_READY
    if workbook_path.is_file():
        _update_cells(workbook_path, row_number, {X_ASSET_STATUS: x_asset_status})
    old_status = str(row.get(TELEGRAM_STATUS) or "")
    old_message_id = str(row.get(TELEGRAM_MESSAGE_ID) or "")
    old_published_at = str(row.get(TELEGRAM_PUBLISHED_AT) or "")
    if old_status in TELEGRAM_SENDING_STATUSES and not force_send:
        raise TelegramSendUncertainError("Telegram 记录仍为发送中；先核对频道结果，不得盲目重发")
    already_sent = old_status in TELEGRAM_SENT_STATUSES or bool(old_message_id)
    if already_sent and not force_send:
        return {
            "status": old_status or "已发送",
            "message_id": old_message_id,
            "published_at": old_published_at,
            "media_type": media_type,
            "media": str(media_path),
            "original_media": str(original_path),
            "x_media": str(x_media),
            "x_status": x_asset_status,
            "telegram_route": telegram_route,
            "caption": str(caption_file),
        }

    command = [
        sys.executable,
        str(TELEGRAM_SCRIPTS["send"]),
        "--video" if media_type == "video" else "--image",
        str(media_path),
        "--caption-file",
        str(caption_file),
    ]
    if send:
        _update_cells(
            workbook_path,
            row_number,
            {TELEGRAM_STATUS: "发送中", TELEGRAM_ERROR: ""},
        )
        command.append("--send")
    try:
        response = _last_json(_run_command(command))
    except PublishPipelineError as exc:
        if send:
            raise TelegramSendUncertainError(
                "Telegram 上传结果无法确认，状态保留为发送中，请先查看目标群再决定是否重试"
            ) from exc
        raise
    status = "已发送" if send else "DRY_RUN"
    message_id = str(response.get("message_id") or "")
    if send and not message_id:
        raise TelegramSendUncertainError(
            "Telegram 返回结果缺少 message_id，状态保留为发送中，请先查看目标群再决定是否重试"
        )
    return {
        "status": status,
        "message_id": message_id,
        "published_at": _now() if send else "",
        "media_type": media_type,
        "media": str(media_path),
        "original_media": str(original_path),
        "x_media": str(x_media),
        "x_status": x_asset_status,
        "telegram_route": telegram_route,
        "caption": str(caption_file),
        "media_size_bytes": media_size,
    }


def _combined_status(
    telegram_status: str,
    x_asset_status: str,
    x_publish_status: str = "",
    typefully_status: str = "",
) -> str:
    telegram_sent = telegram_status in TELEGRAM_SENT_STATUSES
    telegram_failed = telegram_status in FAILED_STATUSES
    telegram_dry = telegram_status == "DRY_RUN"
    x_ready = x_asset_status == X_ASSET_READY
    x_done = (
        x_publish_status in X_DONE_STATUSES
        or typefully_status in TYPEFULLY_DONE_STATUSES
    )

    if telegram_sent and x_done:
        return "COMPLETED"
    if telegram_failed:
        return "PARTIAL_FAILED"
    if telegram_dry and x_ready:
        return "TELEGRAM_DRY_RUN_X_READY"
    if x_done:
        return "X_DONE_TELEGRAM_PENDING"
    if telegram_sent:
        return "TELEGRAM_SENT_X_PENDING"
    if x_ready:
        return "X_READY_TELEGRAM_PENDING"
    if telegram_dry:
        return "DRY_RUN"
    return "PENDING"


def _write_combined_status(workbook_path: Path, row_numbers: Iterable[int]) -> None:
    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    headers = _ensure_columns(sheet, PIPELINE_COLUMNS)
    all_headers = _headers(sheet)
    telegram_status_column = all_headers.get(TELEGRAM_STATUS)
    x_asset_status_column = all_headers.get(X_ASSET_STATUS)
    x_publish_status_column = all_headers.get(X_PUBLISH_STATUS)
    typefully_status_column = all_headers.get(TYPEFULLY_STATUS)
    for row_number in row_numbers:
        telegram_status = str(sheet.cell(row=row_number, column=telegram_status_column).value or "") if telegram_status_column else ""
        x_asset_status = str(sheet.cell(row=row_number, column=x_asset_status_column).value or "") if x_asset_status_column else ""
        x_publish_status = str(sheet.cell(row=row_number, column=x_publish_status_column).value or "") if x_publish_status_column else ""
        typefully_status = str(sheet.cell(row=row_number, column=typefully_status_column).value or "") if typefully_status_column else ""
        errors = []
        for name in (TELEGRAM_ERROR, X_ERROR, TYPEFULLY_ERROR):
            if name in all_headers:
                error = str(sheet.cell(row=row_number, column=all_headers[name]).value or "").strip()
                if error:
                    errors.append(f"{name}: {error}")
        sheet.cell(
            row=row_number,
            column=headers["统一发布状态"],
            value=_combined_status(
                telegram_status,
                x_asset_status,
                x_publish_status,
                typefully_status,
            ),
        )
        sheet.cell(row=row_number, column=headers["统一发布错误信息"], value="；".join(errors))
        sheet.cell(row=row_number, column=headers["统一发布更新时间"], value=_now())
    workbook.save(workbook_path)


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Prepare one bounded video, publish Telegram by Bot API, and prepare Typefully MCP X assets"
    )
    parser.add_argument(
        "--xlsx",
        default=str(FIXED_WORKBOOK_PATH),
        help="Fixed workflow workbook; only output/quark_links.typefully.xlsx is accepted",
    )
    parser.add_argument(
        "--asset-root",
        "--image-root",
        dest="image_root",
        default="output/assets",
        help="Numbered media asset directory",
    )
    parser.add_argument("--quark-db", default="output/quark_state.sqlite3", help="Telegram/Quark metadata SQLite path")
    parser.add_argument("--telegram-cookie-file", default="config/quark-storage.json", help="Quark cookie storage path")
    parser.add_argument(
        "--quark-auto-save-dir",
        default="work/quark-auto-save",
        help="deployed quark-auto-save directory used for complete original downloads",
    )
    parser.add_argument("--limit", type=int, default=1, help="Rows to process in canary mode")
    parser.add_argument("--run-all", action="store_true", help="Process all eligible rows")
    parser.add_argument(
        "--row-number",
        dest="row_numbers",
        action="append",
        type=int,
        help="Process only the specified 1-based Excel row; repeat for a batch",
    )
    parser.add_argument(
        "--only-status",
        default="SHARE_VALIDATED,ASSET_READY,EXPORTED",
        help="Comma-separated processing statuses eligible for preparation",
    )
    parser.add_argument("--telegram-send", action="store_true", help="Actually send Telegram messages")
    parser.add_argument("--force-telegram", action="store_true", help="Send again even when a message ID exists")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    if args.limit < 1:
        raise ValueError("--limit must be greater than zero")

    workbook_path = _resolve_path(args.xlsx)
    if workbook_path != FIXED_WORKBOOK_PATH:
        raise PublishPipelineError(
            "Workflow workbook is fixed at "
            f"{FIXED_WORKBOOK_PATH}; refusing alternate workbook: {workbook_path}"
        )
    image_root = _resolve_path(args.image_root)
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Excel workbook not found: {workbook_path}")

    only_status_values = {value.strip() for value in args.only_status.split(",") if value.strip()}
    if args.row_numbers:
        eligible_rows = _load_candidates(
            workbook_path,
            only_status=only_status_values,
            limit=max(args.row_numbers),
            run_all=True,
            force_reprocess=args.force_telegram,
        )
        missing_rows = [row for row in args.row_numbers if row not in eligible_rows]
        if missing_rows:
            raise PublishPipelineError(
                "Requested Excel rows are not eligible: "
                + ", ".join(str(row) for row in missing_rows)
            )
        row_numbers = list(dict.fromkeys(args.row_numbers))
    else:
        row_numbers = _load_candidates(
            workbook_path,
            only_status=only_status_values,
            limit=args.limit,
            run_all=args.run_all,
            force_reprocess=args.force_telegram,
        )
    if not row_numbers:
        print("publish_candidates=0")
        return 0

    if args.telegram_send:
        _run_command(
            [
                sys.executable,
                str(TELEGRAM_SCRIPTS["send"]),
                "--check",
            ]
        )
        # Fail the entire batch before the first external send if any new row
        # lacks source-backed, byte-identical editorial approval.
        review_workbook = load_workbook(workbook_path)
        review_sheet = review_workbook.active
        review_headers = _headers(review_sheet)
        reviewed = []
        try:
            for row_number in row_numbers:
                row = _row_values(review_sheet, review_headers, row_number)
                serial, title, description, link = _row_identity(row, row_number)
                reviewed.append(
                    validate_review(
                        image_root / f"{serial:04d}",
                        serial=serial, title=title, description=description, link=link,
                    )
                )
            validate_review_batch(reviewed)
        finally:
            review_workbook.close()

    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    _ensure_columns(
        sheet,
        (
            *PIPELINE_COLUMNS,
            *MEDIA_COLUMNS,
            "图片",
            X_ASSET_STATUS,
            X_PUBLISH_STATUS,
            X_ERROR,
            TYPEFULLY_DRAFT_ID,
            TYPEFULLY_STATUS,
            TYPEFULLY_ERROR,
        ),
    )
    image_column = _headers(sheet).get("图片")
    if image_column is None:
        raise PublishPipelineError("无法创建 Excel 图片兜底列")
    workbook.save(workbook_path)
    workbook.close()

    print(
        f"publish_candidates={len(row_numbers)} telegram_send={int(args.telegram_send)} "
        "x_mode=typefully-mcp continue_with=telegram-auto-publish-typefully-mcp"
    )
    for row_number in row_numbers:
        workbook = load_workbook(workbook_path)
        sheet = workbook.active
        row = _row_values(sheet, _headers(sheet), row_number)
        workbook.close()
        try:
            telegram_result = _prepare_telegram_row(
                workbook_path,
                row_number,
                row,
                image_root=image_root,
                quark_db_path=_resolve_path(args.quark_db),
                cookie_file=_resolve_path(args.telegram_cookie_file),
                send=args.telegram_send,
                force_send=args.force_telegram,
                image_column=image_column,
                quark_auto_save_dir=_resolve_path(args.quark_auto_save_dir),
            )
            updates = {
                TELEGRAM_STATUS: telegram_result["status"],
                TELEGRAM_MESSAGE_ID: telegram_result["message_id"],
                TELEGRAM_PUBLISHED_AT: telegram_result["published_at"],
                TELEGRAM_ERROR: "",
                "媒体类型": telegram_result["media_type"],
                "媒体路径": telegram_result["media"],
                "视频": telegram_result["original_media"] if telegram_result["media_type"] == "video" else "",
                "原视频": telegram_result["original_media"] if telegram_result["media_type"] == "video" else "",
                "Telegram视频": telegram_result["original_media"] if telegram_result["media_type"] == "video" else "",
                X_ASSET_STATUS: telegram_result["x_status"],
            }
            _update_cells(workbook_path, row_number, updates)
            print(
                f"telegram row={row_number} status={telegram_result['status']} "
                f"route={telegram_result['telegram_route']} media={telegram_result['media']} "
                f"caption={telegram_result['caption']}"
            )
        except TelegramSendUncertainError as exc:
            error = str(exc)
            _update_cells(workbook_path, row_number, {TELEGRAM_ERROR: error[:1000]})
            print(f"telegram row={row_number} status=发送中 error={error}")
        except CopyReviewError as exc:
            error = str(exc)
            _update_cells(workbook_path, row_number, {TELEGRAM_ERROR: error[:1000]})
            print(f"telegram row={row_number} status=COPY_REVIEW_REQUIRED error={error}")
        except (PublishPipelineError, OSError, ValueError) as exc:
            error = str(exc)
            _update_cells(
                workbook_path,
                row_number,
                {TELEGRAM_STATUS: "FAILED", TELEGRAM_ERROR: error[:1000]},
            )
            print(f"telegram row={row_number} status=FAILED error={error}")
    _write_combined_status(workbook_path, row_numbers)
    print(f"publish_complete workbook={workbook_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
