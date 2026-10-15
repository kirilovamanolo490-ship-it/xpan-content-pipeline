"""知识渠道自动化流程：飞书自动更新 → 收录.md → 人工筛选归档 → Excel 台账。

子命令：

- ingest   ：抓取飞书公开 SSR，把「新增资料」按日期追加到收录.md
- archive  ：读取归档.md，把「新增条目」自动预分类并写入 Excel 台账；
             配合 --quark-run --execute 时自动执行夸克转存 + 新分享 + 匿名验证，
             并把新链接回写 Excel。

人工筛选保持不变：你在 Obsidian 里把选中的条目从 收录.md 剪切到 归档.md，
程序只负责自动化环节，并把结果追加到既有发布主表。

状态持久化：output/knowledge_state.sqlite3（collected / archived 两张表），
断点续跑：已处理的 share_id 不会重复处理。
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sqlite3
import stat
import subprocess
import sys
import tempfile
import time
from contextlib import closing
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Iterable, Mapping
from urllib.parse import urlsplit
from urllib.request import Request, build_opener, HTTPCookieProcessor
from http.cookiejar import CookieJar

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

from .deployment import setting
from .excel_exporter import EXCEL_COLUMNS
from .knowledge_categories import CATEGORIES, classify
from .quark_adapter import QuarkAdapter, QuarkAPIError

# ---------------------------------------------------------------------------
# 常量
# ---------------------------------------------------------------------------

DEFAULT_SOURCE_URL = setting("XPAN_FEISHU_SOURCE_URL")
DEFAULT_SCRAPE_OUT = "work/feishu-source.json"
DEFAULT_COLLECTED_MD = "knowledge/资料渠道一收录.md"
DEFAULT_ARCHIVE_MD = "knowledge/资料渠道一归档.md"
DEFAULT_STATE_DB = "output/knowledge_state.sqlite3"
DEFAULT_LEDGER_XLSX = "output/quark_links.typefully.xlsx"
DEFAULT_COOKIE_FILE = "work/QuarkPanTool/config/cookies.txt"
DEFAULT_TARGET_ROOT = "/推文分享"
DEFAULT_QUARK_URL_FILE = "work/QuarkPanTool/url.txt"
DEFAULT_ARCHIVE_CURSOR = "knowledge/资料渠道一归档游标.md"

_QUARK_URL_RE = re.compile(r"https?://pan\.quark\.cn/s/[A-Za-z0-9_-]+", re.IGNORECASE)
_DATE_HEADER_RE = re.compile(r"^#?\s*📆\s*(\d{2})/(\d{2})")
_ARCHIVE_FOLDER_INDEX_RE = re.compile(r"(?:^|/)(\d+)_")
_ARCHIVE_CURSOR_FIELD_RE = re.compile(r"^\s*-\s*([^:：]+)[：:]\s*(.*?)\s*$")
_ZERO_WIDTH = "\u200b"

LEDGER_COLUMNS = list(EXCEL_COLUMNS)
LEDGER_SHEET_NAME = "处理结果"
DEFAULT_LEDGER_DESCRIPTION = "来源飞书页面已识别，待转存后补充目录说明。"
TERMINAL_ARCHIVE_STATUSES = frozenset({"SHARE_VALIDATED", "SHARED", "SAVED", "BLOCKED"})


# ---------------------------------------------------------------------------
# 数据模型
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class Entry:
    """一条资料：日期分区（MM/DD，可空）+ 标题 + 夸克链接。"""

    date: str
    title: str
    url: str

    @property
    def share_id(self) -> str:
        parts = urlsplit(self.url)
        path_parts = [part for part in parts.path.split("/") if part]
        try:
            index = next(i for i, part in enumerate(path_parts) if part.lower() == "s")
            return path_parts[index + 1] if index + 1 < len(path_parts) else ""
        except StopIteration:
            return ""


@dataclass(frozen=True)
class ArchiveCursor:
    """归档日期游标；扫描起点含当日，便于发现同日后续新增条目。"""

    last_completed_date: str
    scan_from_date: str


def _validate_archive_cursor_date(value: str, field_name: str) -> str:
    value = value.strip()
    match = re.fullmatch(r"(\d{2})/(\d{2})", value)
    if not match:
        raise ValueError(f"归档游标字段{field_name}格式应为 MM/DD：{value or '<empty>'}")
    try:
        datetime(2000, int(match.group(1)), int(match.group(2)))
    except ValueError as exc:
        raise ValueError(f"归档游标字段{field_name}不是有效日期：{value}") from exc
    return value


def read_archive_cursor(path: str | Path) -> ArchiveCursor:
    """读取归档游标；缺失或损坏时失败闭环，避免误处理历史全量。"""
    cursor_path = Path(path)
    if not cursor_path.is_file():
        raise FileNotFoundError(
            f"归档游标不存在：{cursor_path}；请先建立“最后完成日期”和“扫描起点日期（含）”。"
        )
    fields: dict[str, str] = {}
    for line in cursor_path.read_text(encoding="utf-8").splitlines():
        match = _ARCHIVE_CURSOR_FIELD_RE.match(line)
        if match:
            fields[match.group(1).strip()] = match.group(2).strip()
    last_completed = _validate_archive_cursor_date(
        fields.get("最后完成日期", ""), "最后完成日期"
    )
    scan_from = _validate_archive_cursor_date(
        fields.get("扫描起点日期（含）", ""), "扫描起点日期（含）"
    )
    if _date_text_key(scan_from) < _date_text_key(last_completed):
        raise ValueError(
            f"归档游标扫描起点不能早于最后完成日期：{scan_from} < {last_completed}"
        )
    return ArchiveCursor(last_completed_date=last_completed, scan_from_date=scan_from)


def write_archive_cursor(
    path: str | Path,
    cursor: ArchiveCursor,
    *,
    run_at: str,
    run_status: str,
    scope_dates: Iterable[str],
    candidate_count: int,
    processed_count: int,
    pending_dates: Iterable[str],
) -> None:
    """写回人类可读且可被下一次运行读取的归档游标。"""
    scope = ", ".join(scope_dates) or "无"
    pending = ", ".join(pending_dates) or "无"
    cursor_path = Path(path)
    cursor_path.parent.mkdir(parents=True, exist_ok=True)
    cursor_path.write_text(
        "\n".join(
            [
                "# 资料渠道一归档游标",
                "",
                f"- 最后完成日期：{cursor.last_completed_date}",
                f"- 扫描起点日期（含）：{cursor.scan_from_date}",
                f"- 最近一次运行时间：{run_at}",
                f"- 最近一次运行状态：{run_status}",
                f"- 最近一次扫描日期：{scope}",
                f"- 最近一次候选条数：{candidate_count}",
                f"- 最近一次处理条数：{processed_count}",
                f"- 未完成日期：{pending}",
                "",
                "规则：只扫描起点日期及之后的归档条目；已进入终态的条目由 SQLite/Excel 台账跳过。",
                "规则：只有连续日期全部进入终态后，游标才向后滑动；扫描起点保留最后完成日以接住同日新增。",
                "",
            ]
        ),
        encoding="utf-8",
    )


def _clean_title(text: str) -> str:
    text = text.replace(_ZERO_WIDTH, "").strip()
    text = re.sub(r"^#{1,6}\s*", "", text)
    text = re.sub(r"^\s*[-*+]\s+", "", text)
    text = text.replace("**", "").replace("`", "")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _extract_url(line: str) -> str:
    """从普通文本或 markdown 链接 [url](url) 中取夸克 URL。"""
    match = _QUARK_URL_RE.search(line)
    return match.group(0) if match else ""


def _clean_url(url: str) -> str:
    return url.rstrip("，。；、,.!?！？;:：)]}》」』\"'")

# ---------------------------------------------------------------------------
# Markdown 解析（收录.md / 归档.md / Firecrawl 抓取结果共用）
# ---------------------------------------------------------------------------


def _iter_markdown_entries(lines: Iterable[str]) -> Iterable[Entry]:
    """逐行解析 Markdown，避免为大收录池构造整份文本副本。"""
    current_date = ""
    pending_title = ""
    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue
        date_match = _DATE_HEADER_RE.search(line)
        if date_match:
            current_date = f"{date_match.group(1)}/{date_match.group(2)}"
            pending_title = ""
            continue
        url = _extract_url(line)
        if url:
            if not pending_title:
                # 没有标题时退化为 URL 自身
                pending_title = url
            yield Entry(
                date=current_date,
                title=_clean_title(pending_title),
                url=_clean_url(url),
            )
            pending_title = ""
        else:
            # 非链接、非日期头的行视为标题候选
            cleaned = _clean_title(line)
            if cleaned and not cleaned.startswith("["):
                pending_title = cleaned


def _unique_markdown_entries(lines: Iterable[str]) -> tuple[list[Entry], int]:
    """按首次出现保留条目，并返回被压掉的重复 share_id 条数。"""
    unique: dict[str, Entry] = {}
    duplicates = 0
    for entry in _iter_markdown_entries(lines):
        share_id = entry.share_id
        if share_id in unique:
            duplicates += 1
            continue
        unique[share_id] = entry
    return list(unique.values()), duplicates


def parse_markdown(text: str) -> list[Entry]:
    """解析「# 📆 MM/DD → 标题 → 夸克链接」结构并按 share_id 去重。"""
    entries, _duplicates = _unique_markdown_entries(text.splitlines())
    return entries


def _read_markdown_file(path: Path) -> tuple[list[Entry], int]:
    with path.open("r", encoding="utf-8") as handle:
        return _unique_markdown_entries(handle)


def entries_by_share_id(entries: Iterable[Entry]) -> dict[str, Entry]:
    result: dict[str, Entry] = {}
    for entry in entries:
        if entry.share_id:
            result.setdefault(entry.share_id, entry)
    return result


def _ordered_dates(entries: Iterable[Entry]) -> list[str]:
    """按来源出现顺序返回去重后的日期。"""
    dates: list[str] = []
    seen: set[str] = set()
    for entry in entries:
        if entry.date and entry.date not in seen:
            dates.append(entry.date)
            seen.add(entry.date)
    return dates


def _date_key(entry: Entry) -> tuple[int, int]:
    """MM/DD → (月, 日)，解析失败返回 (99, 99) 排到最后。"""
    if not entry.date:
        return (99, 99)
    try:
        month, day = (int(part) for part in entry.date.split("/"))
        return (month, day)
    except ValueError:
        return (99, 99)


def _date_text_key(date: str) -> tuple[int, int]:
    return _date_key(Entry(date=date, title="", url=""))


def _entries_after_collection_watermark(
    source_entries: list[Entry],
    existing_entries: list[Entry],
) -> tuple[list[Entry], str, list[str]]:
    """按收录.md顶部日期选择本次需要检查的来源日期。

    收录.md为空时只取来源最新日期，避免首次运行直接导入历史全量数据。
    已有水位时，取水位之后的所有日期，并包含水位当天，用于发现当天新增链接。
    """
    source_dates = _ordered_dates(source_entries)
    watermark = existing_entries[0].date if existing_entries else ""
    if not source_dates:
        return [], watermark, []

    if not watermark:
        selected_dates = source_dates[:1]
    elif watermark in source_dates:
        # 飞书页面通常按日期倒序；保留到水位当天，且不依赖条目是否已存在。
        end = source_dates.index(watermark) + 1
        selected_dates = source_dates[:end]
    else:
        watermark_key = _date_text_key(watermark)
        selected_dates = [
            date for date in source_dates
            if _date_text_key(date) > watermark_key
        ]

    selected_set = set(selected_dates)
    selected_entries = [entry for entry in source_entries if entry.date in selected_set]
    return selected_entries, watermark, selected_dates


def _render_collection_entry(entry: Entry) -> str:
    """渲染条目；无标题时只写一次链接，避免链接被当成标题再次解析。"""
    title_url = _extract_url(entry.title)
    if title_url and Entry(entry.date, "", title_url).share_id == entry.share_id:
        return entry.url
    return f"{entry.title}\n\n{entry.url}"


def append_entries_to_md(md_path: Path, entries: list[Entry]) -> tuple[int, int]:
    """合并收录条目并压缩重复 share_id，返回（新增条数，移除重复数）。

    读取时逐行处理，仅保留首次出现的 share_id；仅在有新增或重复时重写，
    并通过同目录临时文件原子替换，避免中断时留下半份收录池。
    """
    sections: dict[tuple[int, int], list[Entry]] = {}
    existing_entries: list[Entry] = []
    duplicates_removed = 0
    if md_path.exists():
        existing_entries, duplicates_removed = _read_markdown_file(md_path)
    known_ids: set[str] = set()
    for entry in existing_entries:
        sections.setdefault(_date_key(entry), []).append(entry)
        known_ids.add(entry.share_id)

    inserted = 0
    for entry in entries:
        if entry.share_id in known_ids:
            continue
        key = _date_key(entry)
        section = sections.setdefault(key, [])
        section.append(entry)
        known_ids.add(entry.share_id)
        inserted += 1
    if inserted == 0 and duplicates_removed == 0:
        return 0, 0

    existing_mode = stat.S_IMODE(md_path.stat().st_mode) if md_path.exists() else None
    md_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            newline="\n",
            prefix=f".{md_path.name}.",
            suffix=".tmp",
            dir=md_path.parent,
            delete=False,
        ) as output:
            temporary_path = Path(output.name)
            first_section = True
            for key in sorted(sections, reverse=True):
                if not first_section:
                    output.write("\n")
                first_section = False
                month, day = key
                header = (
                    f"# 📆 {month:02d}/{day:02d}"
                    if key != (99, 99)
                    else "# 📆 未标注日期"
                )
                output.write(header + "\n\n")
                output.write("\n\n".join(_render_collection_entry(e) for e in sections[key]))
                output.write("\n")
            output.flush()
            os.fsync(output.fileno())
        if existing_mode is not None:
            os.chmod(temporary_path, existing_mode)
        os.replace(temporary_path, md_path)
        temporary_path = None
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
    return inserted, duplicates_removed


# ---------------------------------------------------------------------------
# SQLite 状态
# ---------------------------------------------------------------------------


class KnowledgeState:
    def __init__(self, path: str | Path):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._init_db()

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.path)
        connection.row_factory = sqlite3.Row
        return connection

    def _init_db(self) -> None:
        with closing(self._connect()) as connection:
            with connection:
                connection.execute(
                    """
                    CREATE TABLE IF NOT EXISTS collected (
                        share_id TEXT PRIMARY KEY,
                        title TEXT NOT NULL,
                        url TEXT NOT NULL,
                        source_date TEXT,
                        collected_at TEXT NOT NULL
                    )
                    """
                )
                connection.execute(
                    """
                    CREATE TABLE IF NOT EXISTS archived (
                        share_id TEXT PRIMARY KEY,
                        title TEXT NOT NULL,
                        url TEXT NOT NULL,
                        archive_date TEXT,
                        category TEXT,
                        new_share_url TEXT,
                        new_passcode TEXT,
                        save_path TEXT,
                        status TEXT NOT NULL DEFAULT 'EXTRACTED',
                        error TEXT,
                        updated_at TEXT NOT NULL
                    )
                    """
                )

    def mark_collected(self, entry: Entry, collected_at: str) -> None:
        with closing(self._connect()) as connection:
            with connection:
                connection.execute(
                    """
                    INSERT OR IGNORE INTO collected(share_id, title, url, source_date, collected_at)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (entry.share_id, entry.title, entry.url, entry.date, collected_at),
                )

    def collected_ids(self) -> set[str]:
        with closing(self._connect()) as connection:
            return {str(row["share_id"]) for row in connection.execute("SELECT share_id FROM collected")}

    def archived_status(self) -> dict[str, str]:
        with closing(self._connect()) as connection:
            return {
                str(row["share_id"]): str(row["status"])
                for row in connection.execute("SELECT share_id, status FROM archived")
            }

    def archived_error(self) -> dict[str, str]:
        with closing(self._connect()) as connection:
            return {
                str(row["share_id"]): str(row["error"] or "")
                for row in connection.execute("SELECT share_id, error FROM archived")
            }

    def archived_save_paths(self) -> list[str]:
        with closing(self._connect()) as connection:
            return [
                str(row["save_path"])
                for row in connection.execute(
                    "SELECT save_path FROM archived WHERE save_path IS NOT NULL AND save_path != ''"
                )
            ]

    def upsert_archived(
        self,
        entry: Entry,
        *,
        category: str,
        new_share_url: str = "",
        new_passcode: str = "",
        save_path: str = "",
        status: str,
        error: str = "",
        updated_at: str,
    ) -> None:
        with closing(self._connect()) as connection:
            with connection:
                connection.execute(
                    """
                    INSERT INTO archived(
                        share_id, title, url, archive_date, category, new_share_url,
                        new_passcode, save_path, status, error, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(share_id) DO UPDATE SET
                        title=excluded.title, url=excluded.url, archive_date=excluded.archive_date,
                        category=excluded.category, new_share_url=excluded.new_share_url,
                        new_passcode=excluded.new_passcode, save_path=excluded.save_path,
                        status=excluded.status, error=excluded.error, updated_at=excluded.updated_at
                    """,
                    (
                        entry.share_id, entry.title, entry.url, entry.date, category,
                        new_share_url, new_passcode, save_path, status, error, updated_at,
                    ),
                )

    def get(self, share_id: str) -> sqlite3.Row | None:
        with closing(self._connect()) as connection:
            return connection.execute(
                "SELECT * FROM archived WHERE share_id=?", (share_id,)
            ).fetchone()


# ---------------------------------------------------------------------------
# Excel 台账
# ---------------------------------------------------------------------------


def _atomic_save(workbook, output: Path) -> None:
    temporary = output.with_name(f".{output.stem}.{os.getpid()}.tmp{output.suffix}")
    try:
        workbook.save(temporary)
        # Windows may keep the original workbook handle open after save. Close it
        # before replacing the file so later archive rows do not fail on a lock.
        workbook.close()
        for attempt in range(10):
            try:
                os.replace(temporary, output)
                break
            except PermissionError:
                if attempt == 9:
                    raise
                time.sleep(0.5)
    finally:
        workbook.close()
        if temporary.exists():
            temporary.unlink()


def _header_map(sheet) -> dict[str, int]:
    return {
        str(cell.value): cell.column
        for cell in sheet[1]
        if cell.value not in (None, "")
    }


def _ledger_sheet(workbook):
    if LEDGER_SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"Excel 缺少工作表：{LEDGER_SHEET_NAME}")
    return workbook[LEDGER_SHEET_NAME]


def ensure_ledger(xlsx: Path) -> None:
    """验证现有主表结构；归档流程只允许原位追加，不创建替代台账。"""
    if not xlsx.is_file():
        raise FileNotFoundError(f"归档主表不存在：{xlsx}")
    workbook = load_workbook(xlsx, read_only=True)
    try:
        sheet = _ledger_sheet(workbook)
        headers = _header_map(sheet)
        missing = [name for name in LEDGER_COLUMNS if name not in headers]
        if missing:
            raise ValueError(f"归档主表缺少必要列：{', '.join(missing)}")
    finally:
        workbook.close()


def _copy_previous_row_format(sheet, row_number: int) -> None:
    source_row = row_number - 1
    if source_row < 2:
        return
    for column in range(1, sheet.max_column + 1):
        source = sheet.cell(row=source_row, column=column)
        target = sheet.cell(row=row_number, column=column)
        if source.has_style:
            target._style = copy(source._style)
    source_dimension = sheet.row_dimensions[source_row]
    target_dimension = sheet.row_dimensions[row_number]
    target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden
    target_dimension.outlineLevel = source_dimension.outlineLevel


def _extend_auto_filter(sheet, row_number: int) -> None:
    if not sheet.auto_filter.ref:
        return
    min_col, min_row, max_col, _max_row = range_boundaries(sheet.auto_filter.ref)
    sheet.auto_filter.ref = (
        f"{get_column_letter(min_col)}{min_row}:"
        f"{get_column_letter(max_col)}{row_number}"
    )


def _format_bytes(size_bytes: int) -> str:
    value = float(max(0, size_bytes))
    units = ("B", "KB", "MB", "GB", "TB", "PB")
    unit = units[0]
    for candidate in units:
        unit = candidate
        if value < 1024 or candidate == units[-1]:
            break
        value /= 1024
    return f"{int(value)} B" if unit == "B" else f"{value:.2f} {unit}"


def _archive_folder_index(value: str) -> int:
    match = _ARCHIVE_FOLDER_INDEX_RE.search(value.replace("\\", "/"))
    return int(match.group(1)) if match else 0


def _ledger_save_paths(xlsx: Path) -> list[str]:
    if not xlsx.exists():
        return []
    workbook = None
    try:
        workbook = load_workbook(xlsx, read_only=True)
        sheet = _ledger_sheet(workbook)
        save_path_column = _header_map(sheet).get("保存目录")
        if save_path_column is None:
            return []
        return [
            str(row[0].value)
            for row in sheet.iter_rows(min_row=2, min_col=save_path_column, max_col=save_path_column)
            if row[0].value
        ]
    except Exception:
        return []
    finally:
        if workbook is not None:
            workbook.close()


def _ledger_sequence_by_share_id(xlsx: Path) -> dict[str, int]:
    if not xlsx.exists():
        return {}
    workbook = None
    try:
        workbook = load_workbook(xlsx, read_only=True)
        sheet = _ledger_sheet(workbook)
        headers = _header_map(sheet)
        sequence_col = headers.get("序号")
        url_col = headers.get("原夸克链接")
        if sequence_col is None or url_col is None:
            return {}
        sequences: dict[str, int] = {}
        for row in sheet.iter_rows(min_row=2, max_col=max(sequence_col, url_col)):
            values = {column: cell.value for column, cell in enumerate(row, start=1)}
            match = _QUARK_URL_RE.search(str(values.get(url_col) or ""))
            sequence = values.get(sequence_col)
            if not match or not str(sequence or "").isdigit():
                continue
            share_id = Entry("", "", match.group(0)).share_id
            sequences[share_id] = int(sequence)
        return sequences
    except Exception:
        return {}
    finally:
        if workbook is not None:
            workbook.close()


def _next_archive_index(
    state: KnowledgeState,
    xlsx: Path,
    root_items: Iterable[object],
) -> int:
    known_indices = [_archive_folder_index(path) for path in state.archived_save_paths()]
    known_indices.extend(_archive_folder_index(path) for path in _ledger_save_paths(xlsx))
    known_indices.extend(_ledger_sequence_by_share_id(xlsx).values())
    known_indices.extend(
        _archive_folder_index(str(getattr(item, "name", "")))
        for item in root_items
        if getattr(item, "is_dir", False)
    )
    return max(known_indices, default=0) + 1


def existing_ledger_processed_ids(xlsx: Path) -> set[str]:
    """返回台账中已处理（有新链接或终态）的 share_id 集合。

    仅登记（处理状态=EXTRACTED、无新链接）的行不算已处理，
    保证「先登记 Excel、后补转存」的两段式流程可以衔接。
    """
    if not xlsx.exists():
        return set()
    workbook = None
    try:
        workbook = load_workbook(xlsx, read_only=True)
        sheet = _ledger_sheet(workbook)
    except Exception:
        return set()
    try:
        headers = _header_map(sheet)
        url_col = headers.get("原夸克链接")
        if url_col is None:
            return set()
        new_col = headers.get("新夸克链接")
        status_col = headers.get("处理状态")
        processed: set[str] = set()
        for row in sheet.iter_rows(min_row=2, max_col=max(url_col, new_col or url_col, status_col or url_col)):
            values = {column: cell.value for column, cell in enumerate(row, start=1)}
            url = values.get(url_col)
            if not url:
                continue
            match = _QUARK_URL_RE.search(str(url))
            if not match:
                continue
            share_id = Entry("", "", match.group(0)).share_id
            new_link = str(values.get(new_col) or "") if new_col else ""
            status = str(values.get(status_col) or "") if status_col else ""
            if new_link.strip() or status in TERMINAL_ARCHIVE_STATUSES:
                processed.add(share_id)
        return processed
    finally:
        if workbook is not None:
            workbook.close()


def upsert_ledger_row(xlsx: Path, values: Mapping[str, object]) -> int:
    """按原夸克链接 upsert 一行；返回行号。

    已存在的行只更新主表的夸克字段，保留发布列；新行追加到末尾并复制上一行格式。
    """
    ensure_ledger(xlsx)
    workbook = load_workbook(xlsx)
    sheet = _ledger_sheet(workbook)
    headers = _header_map(sheet)
    url_col = headers.get("原夸克链接")
    target_share = Entry("", "", str(values.get("原夸克链接", ""))).share_id
    row_number = None
    for row in sheet.iter_rows(min_row=2, min_col=url_col, max_col=url_col):
        value = row[0].value
        if value:
            match = _QUARK_URL_RE.search(str(value))
            if match and Entry("", "", match.group(0)).share_id == target_share:
                row_number = row[0].row
                break
    effective_values = dict(values)
    if row_number is None:
        row_number = sheet.max_row + 1
        sequence_cells = next(
            sheet.iter_cols(
                min_col=headers["序号"], max_col=headers["序号"], min_row=2
            ),
            (),
        )
        existing_sequences = [
            int(cell.value)
            for cell in sequence_cells
            if str(cell.value or "").isdigit()
        ]
        effective_values.setdefault("序号", max(existing_sequences, default=0) + 1)
        _copy_previous_row_format(sheet, row_number)
        _extend_auto_filter(sheet, row_number)
    for name in LEDGER_COLUMNS:
        if name not in effective_values:
            continue
        sheet.cell(row=row_number, column=headers[name], value=effective_values[name])
    _atomic_save(workbook, xlsx)
    return row_number


def _archive_window_entries(
    entries: Iterable[Entry],
    cursor: ArchiveCursor,
) -> list[Entry]:
    """只保留游标扫描起点（含）之后、按 share_id 去重的归档条目。"""
    unique_entries = entries_by_share_id(entries)
    return [
        entry
        for entry in unique_entries.values()
        if entry.date and _date_text_key(entry.date) >= _date_text_key(cursor.scan_from_date)
    ]


def _load_archive_scope(args: argparse.Namespace) -> tuple[list[Entry], ArchiveCursor | None]:
    md_path = Path(args.md)
    if not md_path.exists():
        print(f"archive_md_missing={md_path}")
        return [], None
    cursor = read_archive_cursor(args.cursor_file)
    entries = parse_markdown(md_path.read_text(encoding="utf-8"))
    return _archive_window_entries(entries, cursor), cursor


def _archive_candidates_from_window(
    window_entries: Iterable[Entry],
    args: argparse.Namespace,
) -> list[Entry]:
    state = KnowledgeState(args.db)
    statuses = state.archived_status()
    ledger_processed = existing_ledger_processed_ids(Path(args.xlsx))

    candidates: list[Entry] = []
    for entry in window_entries:
        share_id = entry.share_id
        status = statuses.get(share_id)
        if status in TERMINAL_ARCHIVE_STATUSES:
            continue
        if status == "FAILED" and not args.retry_failed:
            continue
        if status is None and share_id in ledger_processed:
            # 台账里已有新链接/终态（历史数据），无需重复处理
            continue
        candidates.append(entry)
    return candidates


# ---------------------------------------------------------------------------
# 抓取
# ---------------------------------------------------------------------------


SCRAPE_MAX_ATTEMPTS = 3
SCRAPE_RETRY_DELAY_SECONDS = 2
SCRAPE_TIMEOUT_SECONDS = 30


def _feishu_block_text(block: Mapping[str, object]) -> str:
    data = block.get("data")
    if not isinstance(data, Mapping):
        return ""
    text = data.get("text")
    if not isinstance(text, Mapping):
        return ""
    attributed = text.get("initialAttributedTexts")
    if not isinstance(attributed, Mapping):
        return ""
    fragments = attributed.get("text")
    if not isinstance(fragments, Mapping):
        return ""
    return "".join(
        str(fragments[key])
        for key in sorted(fragments, key=lambda value: int(value))
    ).encode("utf-8", "replace").decode("utf-8")


def _extract_feishu_markdown(html: str) -> str:
    """从飞书公开页的 SSR 数据中按文档顺序提取正文。"""
    marker = "clientVars: Object("
    start = html.find(marker)
    if start < 0:
        raise ValueError("missing_client_vars")
    try:
        payload, _ = json.JSONDecoder().raw_decode(html[start + len(marker) :])
        blocks = payload["data"]["block_map"]
    except (KeyError, TypeError, ValueError, json.JSONDecodeError) as exc:
        raise ValueError(f"invalid_client_vars={type(exc).__name__}") from exc
    if not isinstance(blocks, Mapping):
        raise ValueError("invalid_block_map")

    page = next(
        (
            block.get("data")
            for block in blocks.values()
            if isinstance(block, Mapping)
            and isinstance(block.get("data"), Mapping)
            and block["data"].get("type") == "page"
        ),
        None,
    )
    if not isinstance(page, Mapping) or not isinstance(page.get("children"), list):
        raise ValueError("missing_page_children")

    lines: list[str] = []
    for block_id in page["children"]:
        block = blocks.get(block_id)
        if not isinstance(block, Mapping):
            continue
        content = _feishu_block_text(block)
        if not content:
            continue
        data = block.get("data")
        block_type = data.get("type", "") if isinstance(data, Mapping) else ""
        lines.append(f"# {content}" if str(block_type).startswith("heading") else content)
    return "\n".join(lines)


def _fetch_feishu_html(source_url: str) -> str:
    """Fetch the public SSR page with the same verified-TLS client on every OS."""
    request = Request(source_url, headers={
        "User-Agent": "Mozilla/5.0", "Accept": "text/html", "Accept-Encoding": "identity",
    })
    opener = build_opener(HTTPCookieProcessor(CookieJar()))
    with opener.open(request, timeout=SCRAPE_TIMEOUT_SECONDS) as response:
        return response.read().decode(response.headers.get_content_charset() or "utf-8")


def _scrape_output_status(scrape_file: Path) -> tuple[str, int, list[str]]:
    """校验抓取结果，返回 (状态, 可解析链接数, 日期分区)。"""
    if not scrape_file.is_file():
        return "output_missing", 0, []
    try:
        markdown = load_scrape_markdown(scrape_file)
    except (OSError, UnicodeError) as exc:
        return f"output_read_error={type(exc).__name__}", 0, []
    if not markdown.strip():
        return "empty_markdown", 0, []

    entries = parse_markdown(markdown)
    parsed_count = len(entries_by_share_id(entries))
    source_dates = _ordered_dates(entries)
    if parsed_count == 0:
        return "no_quark_links", 0, source_dates
    return "ready", parsed_count, source_dates


def scrape_feishu(
    source_url: str,
    out_path: Path,
) -> None:
    """直连固定飞书公开页，并重试直到提取结果可用。"""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    print(
        f"scrape_start url={source_url} out={out_path} "
        f"max_attempts={SCRAPE_MAX_ATTEMPTS}"
    )
    last_reason = "unknown"
    for attempt in range(1, SCRAPE_MAX_ATTEMPTS + 1):
        # Never accept a previous run's valid file when this scrape failed.
        try:
            out_path.unlink(missing_ok=True)
        except OSError as exc:
            raise RuntimeError(f"无法清理旧抓取结果 {out_path}: {exc}") from exc

        print(f"scrape_attempt={attempt}/{SCRAPE_MAX_ATTEMPTS}")
        try:
            out_path.write_text(_extract_feishu_markdown(_fetch_feishu_html(source_url)), encoding="utf-8")
        except (OSError, RuntimeError, ValueError, subprocess.TimeoutExpired) as exc:
            last_reason = f"native_fetch_error={type(exc).__name__}: {str(exc)[:500]}"
        else:
            status, parsed_count, source_dates = _scrape_output_status(out_path)
            last_reason = status
            if status == "ready":
                latest = source_dates[0] if source_dates else "<none>"
                print(f"scrape_ready attempt={attempt} parsed={parsed_count} source_latest={latest}")
                return

        print(f"scrape_attempt_failed attempt={attempt} reason={last_reason}")
        if attempt < SCRAPE_MAX_ATTEMPTS:
            print(
                f"scrape_retry next_attempt={attempt + 1} "
                f"delay_seconds={SCRAPE_RETRY_DELAY_SECONDS}"
            )
            time.sleep(SCRAPE_RETRY_DELAY_SECONDS)

    raise RuntimeError(
        f"飞书原生抓取在 {SCRAPE_MAX_ATTEMPTS} 次尝试后仍未通过校验: {last_reason}"
    )


def load_scrape_markdown(scrape_file: Path) -> str:
    """读取 SSR 提取的 Markdown；兼容旧版 Firecrawl markdown+links JSON。"""
    content = scrape_file.read_text(encoding="utf-8")
    try:
        data = json.loads(content)
    except json.JSONDecodeError:
        return content
    if isinstance(data, dict):
        markdown = data.get("markdown") or ""
        if markdown:
            return markdown
        links = data.get("links") or []
        if isinstance(links, list):
            return "\n".join(str(link) for link in links)
    if isinstance(data, list):
        return "\n".join(
            str(item.get("markdown") or item.get("url") or "")
            for item in data
            if isinstance(item, dict)
        )
    return ""


# ---------------------------------------------------------------------------
# 子命令：ingest
# ---------------------------------------------------------------------------


def _date_age_days(entry: Entry, today: datetime) -> int | None:
    """返回条目日期距今天的天数；无法解析日期时返回 None。"""
    if not entry.date:
        return None
    try:
        month, day = (int(part) for part in entry.date.split("/"))
    except ValueError:
        return None
    candidate = today.replace(month=month, day=day)
    if candidate > today:
        candidate = candidate.replace(year=today.year - 1)
    return (today - candidate).days


def _cmd_ingest(args: argparse.Namespace) -> int:
    scrape_file = Path(args.scrape_file)
    if args.scrape:
        if not args.source_url:
            raise ValueError("请配置 XPAN_FEISHU_SOURCE_URL 或传入 --source-url")
        scrape_feishu(args.source_url, scrape_file)
    elif not scrape_file.exists():
        print(f"scrape_file_missing={scrape_file} 请先运行抓取或去掉 --scrape")
        return 2

    markdown = load_scrape_markdown(scrape_file)
    if not markdown.strip():
        print(f"scrape_empty={scrape_file} 抓取结果中没有可解析内容")
        return 2

    fresh = parse_markdown(markdown)
    fresh_by_id = entries_by_share_id(fresh)
    print(f"scrape_parsed={len(fresh_by_id)} url={args.source_url}")
    md_path = Path(args.md)
    existing_entries: list[Entry] = []
    existing_by_id: dict[str, Entry] = {}
    if md_path.exists():
        existing_entries, _duplicates = _read_markdown_file(md_path)
        existing_by_id = entries_by_share_id(existing_entries)
    state = KnowledgeState(args.db)
    known_ids = existing_by_id.keys() | state.collected_ids()

    # 默认按收录.md顶部日期做动态水位；显式传入旧筛选参数时保留旧行为，
    # 便于补跑指定日期范围。
    use_legacy_window = args.since_date is not None or args.since_days is not None
    watermark = ""
    selected_dates: list[str] = []
    if use_legacy_window:
        candidate_entries = fresh
        selection_label = "legacy"
    else:
        candidate_entries, watermark, selected_dates = _entries_after_collection_watermark(
            fresh, existing_entries
        )
        selection_label = "collection_watermark"
    candidate_by_id = entries_by_share_id(candidate_entries)
    print(
        f"selection={selection_label} watermark={watermark or '<empty>'} "
        f"source_dates={','.join(_ordered_dates(fresh)[:20]) or '<none>'} "
        f"selected_dates={','.join(selected_dates) or '<all>'}"
    )

    today = datetime.now()
    since_threshold: int | None = None
    if args.since_date:
        try:
            month, day = (int(part) for part in args.since_date.split("/"))
            since = today.replace(month=month, day=day)
            if since > today:
                since = since.replace(year=today.year - 1)
            since_threshold = (today - since).days
        except ValueError:
            print(f"since_date 格式应为 MM/DD：{args.since_date}")
            return 2
    elif args.since_days is not None and args.since_days > 0:
        since_threshold = args.since_days

    new_entries: list[Entry] = []
    for share_id, entry in candidate_by_id.items():
        if share_id in known_ids:
            continue
        if since_threshold is not None:
            age = _date_age_days(entry, today)
            if age is None or age > since_threshold:
                continue
        new_entries.append(entry)
        known_ids.add(share_id)

    window = args.since_date or (
        f"最近{args.since_days}天"
        if args.since_days is not None and args.since_days > 0
        else selection_label
    )
    print(f"existing={len(existing_by_id)} new_in_{window}={len(new_entries)}")
    if args.dry_run:
        for entry in new_entries[:30]:
            print(f"  new date={entry.date} title={entry.title} url={entry.url}")
        if len(new_entries) > 30:
            print(f"  … 还有 {len(new_entries) - 30} 条")
        return 0

    now = datetime.now().astimezone().isoformat(timespec="seconds")
    inserted, compacted = append_entries_to_md(md_path, new_entries)
    for entry in new_entries:
        state.mark_collected(entry, now)
    if compacted:
        print(f"collection_md_compacted duplicate_share_ids_removed={compacted}")
    print(f"ingest_done appended={inserted} md={md_path}")
    return 0


# ---------------------------------------------------------------------------
# 子命令：archive
# ---------------------------------------------------------------------------


def _archive_candidates(args: argparse.Namespace) -> list[Entry]:
    window_entries, _cursor = _load_archive_scope(args)
    return _archive_candidates_from_window(window_entries, args)


def _advance_archive_cursor(
    window_entries: Iterable[Entry],
    state: KnowledgeState,
    xlsx: Path,
    cursor: ArchiveCursor,
) -> tuple[ArchiveCursor, list[str]]:
    """只越过全部进入终态的连续日期，并保留最后完成日作为重扫起点。"""
    statuses = state.archived_status()
    ledger_processed = existing_ledger_processed_ids(xlsx)
    grouped: dict[str, list[Entry]] = {}
    for entry in window_entries:
        if entry.date:
            grouped.setdefault(entry.date, []).append(entry)

    completed_date = cursor.last_completed_date
    pending_dates: list[str] = []
    for date in sorted(grouped, key=_date_text_key):
        complete = all(
            statuses.get(entry.share_id) in TERMINAL_ARCHIVE_STATUSES
            or entry.share_id in ledger_processed
            for entry in grouped[date]
        )
        if complete:
            completed_date = date
        else:
            pending_dates.append(date)
            break

    return (
        ArchiveCursor(
            last_completed_date=completed_date,
            scan_from_date=completed_date,
        ),
        pending_dates,
    )


def _run_quark_for_entry(
    adapter: QuarkAdapter,
    entry: Entry,
    index: int,
    args: argparse.Namespace,
) -> tuple[str, str, str, int, str, str]:
    """返回 (status, new_share_url, new_passcode, file_count, size_display, error)。"""
    adapter.inspect_share(entry.url)
    safe_title = re.sub(r'[\\/:*?"<>|]+', "_", entry.title or "资料").strip()[:80] or "资料"
    target = f"{args.target_root.rstrip('/')}/{index:04d}_{safe_title}"
    folder_id = adapter.ensure_folder(target)
    saved_items = tuple(adapter.list_folder(folder_id))
    if not saved_items:
        saved_items = adapter.save_share(entry.url, folder_id).items
    if not saved_items:
        return "BLOCKED", "", "", 0, "", "转存任务完成，但未发现新的目标文件项；未自动新建分享。"

    file_count = len(saved_items)
    size_display = ""
    size_error = ""
    try:
        usage = adapter.calculate_usage(saved_items)
        size_display = _format_bytes(usage.total_bytes)
    except (QuarkAPIError, OSError, ValueError) as exc:
        size_error = f"文件大小获取失败：{str(exc)[:400]}"

    shares: list[str] = []
    passcodes: list[str] = []
    all_valid = True
    for item in saved_items:
        share = adapter.create_share(
            item,
            expiry=args.share_expiry,
            password_mode=args.share_password_mode,
            password=args.share_password,
        )
        valid = adapter.validate_share(share)
        all_valid = all_valid and valid
        shares.append(share.share_url)
        if share.passcode:
            passcodes.append(share.passcode)
    status = "SHARE_VALIDATED" if all_valid else "SHARED"
    errors = [size_error] if size_error else []
    if not all_valid:
        errors.append("新分享已创建，但公开页面可达性校验未通过。")
    return status, "\n".join(shares), "\n".join(passcodes), file_count, size_display, "；".join(errors)


def _cmd_archive(args: argparse.Namespace) -> int:
    window_entries, cursor = _load_archive_scope(args)
    if cursor is None:
        print(f"archive_candidates=0 md={args.md}")
        return 0

    scope_dates = _ordered_dates(window_entries)
    print(
        f"archive_cursor last_completed={cursor.last_completed_date} "
        f"scan_from={cursor.scan_from_date} scope_dates={','.join(scope_dates) or '<none>'}"
    )
    candidates = _archive_candidates_from_window(window_entries, args)
    print(f"archive_candidates={len(candidates)} md={args.md}")

    for entry in candidates:
        category, hits = classify(entry.title)
        print(f"  candidate date={entry.date} category={category} hits={hits} title={entry.title[:60]} url={entry.url}")

    if args.dry_run:
        return 0

    state = KnowledgeState(args.db)
    xlsx = Path(args.xlsx)
    now = datetime.now().astimezone().isoformat(timespec="seconds")

    if not candidates:
        if window_entries:
            ensure_ledger(xlsx)
            next_cursor, pending_dates = _advance_archive_cursor(
                window_entries, state, xlsx, cursor
            )
        else:
            next_cursor, pending_dates = cursor, []
        run_status = (
            "PARTIAL"
            if pending_dates
            else ("NO_NEW_DATE" if not window_entries else "COMPLETED")
        )
        write_archive_cursor(
            args.cursor_file,
            next_cursor,
            run_at=now,
            run_status=run_status,
            scope_dates=scope_dates,
            candidate_count=0,
            processed_count=0,
            pending_dates=pending_dates,
        )
        print(
            f"archive_cursor_updated last_completed={next_cursor.last_completed_date} "
            f"scan_from={next_cursor.scan_from_date} status={run_status} "
            f"pending={','.join(pending_dates) or '<none>'}"
        )
        print(f"archive_done processed=0 xlsx={xlsx}")
        return 0

    ensure_ledger(xlsx)
    ledger_sequences = _ledger_sequence_by_share_id(xlsx)

    adapter: QuarkAdapter | None = None
    if args.quark_run:
        if not args.execute:
            print("未执行转存；需要 --quark-run 配合 --execute，建议先用默认 limit=1 验证。")
            return 1
        adapter = QuarkAdapter(cookie_file=args.cookie_file)
        nickname = adapter.whoami()
        print(f"quark_user={nickname}")
        target_root_id = adapter.ensure_folder(args.target_root)
        next_archive_index = _next_archive_index(
            state, xlsx, adapter.list_folder(target_root_id)
        )
        print(f"archive_next_folder_index={next_archive_index}")
    else:
        next_archive_index = 1

    candidate_count = len(candidates)
    if args.limit > 0:
        candidates = candidates[: args.limit]

    processed_count = 0
    for entry in candidates:
        index = ledger_sequences.get(entry.share_id)
        if index is None:
            index = next_archive_index
            next_archive_index += 1
        category, _hits = classify(entry.title)
        status = "EXTRACTED"
        new_url = ""
        new_passcode = ""
        file_count: int | str = ""
        size_display = ""
        save_path = ""
        error = ""
        if adapter is not None:
            try:
                status, new_url, new_passcode, file_count, size_display, error = _run_quark_for_entry(
                    adapter, entry, index, args
                )
                safe_title = re.sub(r'[\\/:*?"<>|]+', "_", entry.title or "资料").strip()[:80] or "资料"
                save_path = f"{args.target_root.rstrip('/')}/{index:04d}_{safe_title}"
            except (QuarkAPIError, OSError, ValueError) as exc:
                status = "FAILED"
                error = str(exc)[:500]
                print(f"failed={entry.url} error={exc}")

        values = {
            "序号": index,
            "来源飞书链接": args.source_url,
            "资料标题": entry.title,
            "原夸克链接": entry.url,
            "原提取码": "",
            "新夸克链接": new_url,
            "新提取码": new_passcode,
            "保存目录": save_path,
            "文件数量": file_count,
            "文件类型": "",
            "总大小": size_display,
            "文字介绍": DEFAULT_LEDGER_DESCRIPTION,
            "处理状态": status,
            "错误信息": error,
            "更新时间": now,
        }
        row = upsert_ledger_row(xlsx, values)
        state.upsert_archived(
            entry,
            category=category,
            new_share_url=new_url,
            new_passcode=new_passcode,
            save_path=save_path,
            status=status,
            error=error,
            updated_at=now,
        )
        processed_count += 1
        print(f"row={row} status={status} url={entry.url}")

    next_cursor, pending_dates = _advance_archive_cursor(
        window_entries, state, xlsx, cursor
    )
    run_status = "PARTIAL" if pending_dates else "COMPLETED"
    write_archive_cursor(
        args.cursor_file,
        next_cursor,
        run_at=now,
        run_status=run_status,
        scope_dates=scope_dates,
        candidate_count=candidate_count,
        processed_count=processed_count,
        pending_dates=pending_dates,
    )
    print(
        f"archive_cursor_updated last_completed={next_cursor.last_completed_date} "
        f"scan_from={next_cursor.scan_from_date} status={run_status} "
        f"pending={','.join(pending_dates) or '<none>'}"
    )
    print(f"archive_done processed={processed_count} xlsx={xlsx}")
    return 0


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="知识渠道自动化：飞书增量收录 → 收录.md → 归档.md → 夸克转存 → Excel 台账"
    )
    sub = parser.add_subparsers(dest="command", required=True)

    ingest = sub.add_parser("ingest", help="抓取飞书页面并增量追加到收录.md")
    ingest.add_argument("--source-url", default=DEFAULT_SOURCE_URL)
    ingest.add_argument("--scrape-file", default=DEFAULT_SCRAPE_OUT)
    ingest.add_argument("--md", default=DEFAULT_COLLECTED_MD)
    ingest.add_argument("--db", default=DEFAULT_STATE_DB)
    ingest.add_argument("--scrape", action="store_true", help="直连抓取飞书页面（默认只读已有缓存文件）")
    ingest.add_argument(
        "--since-date", default=None, metavar="MM/DD",
        help="兼容模式：只收录该日期（含）之后的条目；不传时按收录.md顶部日期增量收录",
    )
    ingest.add_argument(
        "--since-days", type=int, default=None,
        help="兼容模式：只收录最近 N 天内的条目（0=不限制；since-date 优先）",
    )
    ingest.add_argument("--dry-run", action="store_true", help="只打印新增条目，不写文件")
    ingest.set_defaults(func=_cmd_ingest)

    archive = sub.add_parser("archive", help="读取归档.md，预分类并写入 Excel；可自动转存/新分享")
    archive.add_argument("--md", default=DEFAULT_ARCHIVE_MD)
    archive.add_argument("--xlsx", default=DEFAULT_LEDGER_XLSX)
    archive.add_argument("--cursor-file", default=DEFAULT_ARCHIVE_CURSOR)
    archive.add_argument("--source-url", default=DEFAULT_SOURCE_URL)
    archive.add_argument("--db", default=DEFAULT_STATE_DB)
    archive.add_argument("--cookie-file", default=DEFAULT_COOKIE_FILE)
    archive.add_argument("--target-root", default=DEFAULT_TARGET_ROOT)
    archive.add_argument("--share-expiry", choices=("1d", "7d", "30d", "permanent"), default="permanent")
    archive.add_argument("--share-password-mode", choices=("none", "generated", "fixed"), default="none")
    archive.add_argument("--share-password", default=None)
    archive.add_argument("--quark-run", action="store_true", help="执行夸克转存+新分享（需配合 --execute）")
    archive.add_argument("--execute", action="store_true", help="真正写夸克网盘")
    archive.add_argument("--limit", type=int, default=1, help="本次最多处理的条数（canary 模式）")
    archive.add_argument("--run-all", action="store_true", help="处理全部候选条目")
    archive.add_argument("--retry-failed", action="store_true", help="重试状态为 FAILED 的条目")
    archive.add_argument("--dry-run", action="store_true", help="只打印候选与预分类，不写任何文件")
    archive.set_defaults(func=_cmd_archive)
    return parser


def _configure_output_encoding() -> None:
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure is None:
            continue
        try:
            reconfigure(encoding="utf-8", errors="backslashreplace")
        except (OSError, ValueError):
            continue


def main(argv: list[str] | None = None) -> int:
    _configure_output_encoding()
    args = _parser().parse_args(argv)
    if getattr(args, "run_all", False):
        args.limit = 0
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
