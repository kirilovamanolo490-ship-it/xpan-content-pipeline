from __future__ import annotations

import argparse
import html
import json
import os
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import load_workbook

from .telegram_formatter import TELEGRAM_CHANNEL_URL
from .models import QuarkLink
from .state import StateStore
from .typefully_publisher import TypefullyAPIError, TypefullyClient
from .video_quality import inspect_publish_video


TEXT_COLUMNS = ("X文案", "X正文", "推文文案", "文案")
TEXT_PATH_COLUMNS = ("X文案路径", "X文案文件", "文案路径", "文案文件")
TITLE_COLUMNS = ("资料标题", "标题", "名称")
DESCRIPTION_COLUMNS = ("文字介绍", "简介", "描述")
LINK_COLUMNS = ("新夸克链接", "新分享链接", "分享链接", "链接")
PASSCODE_COLUMNS = ("新提取码", "提取码", "密码")
IMAGE_COLUMNS = ("图片", "图片路径", "图片文件", "image_path", "image")
VIDEO_COLUMNS = ("视频", "视频路径", "视频文件", "video_path", "video")
MEDIA_COLUMNS = ("媒体路径", "媒体文件", "Typefully媒体路径")
PUBLISH_MEDIA_EXTENSIONS = {".mp4", ".mov", ".jpg", ".jpeg", ".png", ".webp", ".gif"}
TYPEFULLY_VIDEO_EXTENSIONS = {".mp4", ".mov"}
TYPEFULLY_MAX_VIDEO_BYTES = 512 * 1024 * 1024
TYPEFULLY_MAX_VIDEO_SECONDS = 600.0
STATUS_COLUMNS = ("处理状态", "状态", "status")

KNOWLEDGE_GROUP_URL = TELEGRAM_CHANNEL_URL
KNOWLEDGE_GROUP_LINE = f"知识分享群：{KNOWLEDGE_GROUP_URL}"

TYPEFULLY_COLUMNS = (
    "X文案",
    "Typefully媒体路径",
    "Typefully图片路径",
    "Typefully草稿ID",
    "Typefully草稿链接",
    "Typefully状态",
    "Typefully错误信息",
    "Typefully更新时间",
)


def _load_dotenv(path: str | Path | None = None) -> None:
    """Load simple KEY=value entries without overriding explicit environment values."""

    candidates = [Path(path)] if path else [
        Path.cwd() / ".env",
        Path(__file__).resolve().parents[2] / ".env",
    ]
    for dotenv_path in candidates:
        if not dotenv_path.is_file():
            continue
        try:
            lines = dotenv_path.read_text(encoding="utf-8-sig").splitlines()
        except (OSError, UnicodeError):
            return
        for raw_line in lines:
            line = raw_line.strip()
            if not line or line.startswith("#"):
                continue
            if line.startswith("export "):
                line = line[7:].lstrip()
            key, separator, value = line.partition("=")
            if not separator or not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", key.strip()):
                continue
            value = value.strip()
            if len(value) >= 2 and value[0] == value[-1] and value[0] in {"'", '"'}:
                value = value[1:-1]
            os.environ.setdefault(key.strip(), value)
        return


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _value(row: Mapping[str, object], names: Iterable[str]) -> str:
    for name in names:
        value = row.get(name)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def _read_text_file(path: Path) -> str:
    if not path.is_file():
        return ""
    try:
        return path.read_text(encoding="utf-8-sig").strip()
    except (OSError, UnicodeError):
        return ""


def _read_telegram_caption(path: Path) -> str:
    text = _read_text_file(path)
    if not text:
        return ""
    lines = text.splitlines()
    while lines and not lines[-1].strip():
        lines.pop()
    if lines and "<a " in lines[-1].lower():
        lines.pop()
    text = "\n".join(lines)
    return html.unescape(re.sub(r"<[^>]+>", "", text)).strip()


def prepare_x_post_text(text: str) -> str:
    """Keep standard Chinese hashtags for Typefully/X parsing."""

    # X derives hashtag entities from the plain text. Hidden boundary
    # characters make Typefully treat the visible Chinese tags as ordinary
    # text, so keep a normal space before each '#' and no hidden characters.
    value = text.replace("\u200c", "").strip()

    if KNOWLEDGE_GROUP_URL not in value:
        lines = value.splitlines()
        hashtag_index = next(
            (index for index, line in enumerate(lines) if line.strip().startswith("#")),
            None,
        )
        if hashtag_index is None:
            value = f"{value}\n\n{KNOWLEDGE_GROUP_LINE}" if value else KNOWLEDGE_GROUP_LINE
        else:
            before = lines[:hashtag_index]
            while before and not before[-1].strip():
                before.pop()
            after = lines[hashtag_index:]
            value = "\n".join(before) + f"\n\n{KNOWLEDGE_GROUP_LINE}\n\n" + "\n".join(after)
    return value.strip()


def _resolve_asset_root(
    asset_root: str | Path | None,
    *,
    workbook_path: str | Path | None = None,
) -> Path | None:
    if not asset_root:
        return None
    candidate = Path(asset_root)
    if candidate.is_absolute():
        return candidate
    resolved = Path.cwd() / candidate
    if resolved.exists() or workbook_path is None:
        return resolved
    return Path(workbook_path).resolve().parent / candidate


def build_post_text(
    row: Mapping[str, object],
    *,
    data_index: int | None = None,
    asset_root: str | Path | None = "output/assets",
    workbook_path: str | Path | None = None,
) -> str:
    """Reuse the prepared X copy, then Telegram sidecars, before old-sheet fallback."""

    prepared = _value(row, TEXT_COLUMNS)
    if prepared:
        return prepare_x_post_text(prepared)

    if workbook_path:
        for value in (_value(row, TEXT_PATH_COLUMNS),):
            for candidate in _candidate_paths(
                value,
                workbook_path=Path(workbook_path).resolve(),
                image_root=_resolve_asset_root(asset_root, workbook_path=workbook_path),
            ):
                text = _read_text_file(candidate)
                if text:
                    return prepare_x_post_text(text)

    root = _resolve_asset_root(asset_root, workbook_path=workbook_path)
    if root and data_index is not None:
        asset_dir = root / f"{data_index:04d}"
        for filename in ("x.txt", "telegram.txt"):
            text = _read_text_file(asset_dir / filename)
            if text:
                return prepare_x_post_text(text)
        text = _read_telegram_caption(asset_dir / "telegram-caption.html")
        if text:
            return prepare_x_post_text(text)

    parts: list[str] = []
    title = _value(row, TITLE_COLUMNS)
    description = _value(row, DESCRIPTION_COLUMNS)
    link = _value(row, LINK_COLUMNS)
    passcode = _value(row, PASSCODE_COLUMNS)
    if title:
        parts.append(title)
    if description:
        parts.append(description)
    if link:
        parts.append(f"🔗 {link}")
    if passcode:
        parts.append(f"提取码：{passcode}")
    return prepare_x_post_text("\n\n".join(parts))


def _candidate_paths(value: str, *, workbook_path: Path, image_root: Path | None) -> Iterable[Path]:
    if value:
        path = Path(value)
        if path.is_absolute():
            yield path
        else:
            yield workbook_path.parent / path
            yield Path.cwd() / path
            if image_root:
                yield image_root / path


def resolve_image_path(
    row: Mapping[str, object],
    *,
    data_index: int,
    workbook_path: str | Path,
    image_root: str | Path | None = "output/assets",
) -> Path | None:
    """Find an existing local image without extracting Excel's embedded thumbnail."""

    workbook = Path(workbook_path).resolve()
    root = _resolve_asset_root(image_root, workbook_path=workbook)

    for value in (_value(row, IMAGE_COLUMNS), _value(row, ("Typefully图片路径",))):
        for candidate in _candidate_paths(value, workbook_path=workbook, image_root=root):
            if candidate.is_file():
                return candidate.resolve()

    if root:
        numbered_dir = root / f"{data_index:04d}"
        if numbered_dir.is_dir():
            preferred = (
                "cover-v1-information-card.png",
                "cover.png",
                "cover.jpg",
                "cover.jpeg",
                "cover.webp",
            )
            for filename in preferred:
                path = numbered_dir / filename
                if path.is_file():
                    return path.resolve()
            for path in sorted(numbered_dir.iterdir()):
                if path.is_file() and path.suffix.lower() in {".jpg", ".jpeg", ".png", ".webp", ".gif"}:
                    return path.resolve()
    return None


def resolve_media_path(
    row: Mapping[str, object],
    *,
    data_index: int,
    workbook_path: str | Path,
    asset_root: str | Path | None = "output/assets",
) -> Path | None:
    """Resolve the preferred video, or the approved fallback image when no video exists."""

    workbook = Path(workbook_path).resolve()
    root = _resolve_asset_root(asset_root, workbook_path=workbook)
    values = (
        _value(row, ("Typefully媒体路径",)),
        _value(row, VIDEO_COLUMNS),
        _value(row, MEDIA_COLUMNS),
        _value(row, IMAGE_COLUMNS),
        _value(row, ("Typefully图片路径",)),
    )
    for value in values:
        for candidate in _candidate_paths(value, workbook_path=workbook, image_root=root):
            if candidate.is_file() and candidate.suffix.lower() in PUBLISH_MEDIA_EXTENSIONS:
                return candidate.resolve()

    if root:
        numbered_dir = root / f"{data_index:04d}"
        if numbered_dir.is_dir():
            preferred = (
                "project-video-x.mp4",
                "cover-v1-information-card.png",
                "cover.png",
                "cover.jpg",
                "cover.jpeg",
                "cover.webp",
            )
            for filename in preferred:
                path = numbered_dir / filename
                if path.is_file():
                    return path.resolve()
    return None


def validate_typefully_media(media: Path | None) -> str | None:
    """Return a channel-specific validation error without blocking Telegram."""

    if media is None or media.suffix.lower() not in TYPEFULLY_VIDEO_EXTENSIONS:
        return None
    report = inspect_publish_video(
        media,
        max_size_bytes=TYPEFULLY_MAX_VIDEO_BYTES,
        max_duration_seconds=TYPEFULLY_MAX_VIDEO_SECONDS,
    )
    if report.passed:
        return None
    return "Typefully 视频不符合 X 发布限制：" + "；".join(report.issues)


def _headers(sheet) -> dict[str, int]:
    result: dict[str, int] = {}
    for index, cell in enumerate(sheet[1], start=1):
        if cell.value is not None:
            result[str(cell.value)] = index
    return result


def _ensure_columns(sheet, names: Iterable[str]) -> dict[str, int]:
    headers = _headers(sheet)
    for name in names:
        if name in headers:
            continue
        column = sheet.max_column + 1
        sheet.cell(row=1, column=column, value=name)
        headers[name] = column
    return headers


def _row_values(sheet, headers: Mapping[str, int], row_number: int) -> dict[str, object]:
    return {name: sheet.cell(row=row_number, column=column).value for name, column in headers.items()}


def _job_key(row: Mapping[str, object]) -> str | None:
    source = _value(row, ("原夸克链接", "原始夸克链接", "新夸克链接", "新分享链接"))
    if not source:
        return None
    try:
        return QuarkLink(original_url=source, normalized_url=source).key
    except ValueError:
        return None


def _record_state(
    store: StateStore | None,
    row: Mapping[str, object],
    *,
    status: str,
    metadata: Mapping[str, object],
    error: str | None = None,
) -> None:
    if store is None:
        return
    job_key = _job_key(row)
    if not job_key:
        return
    store.update_metadata(job_key, {"typefully": dict(metadata)}, status=status, error=error)


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Upload local prepared posts to Typefully as X drafts")
    parser.add_argument("--xlsx", default=None, help="Existing prepared Excel workbook")
    parser.add_argument("--api-key", default=os.getenv("TYPEFULLY_API_KEY", ""), help="Typefully API key or TYPEFULLY_API_KEY")
    parser.add_argument(
        "--social-set-id",
        default=os.getenv("TYPEFULLY_SOCIAL_SET_ID", ""),
        help="Typefully Social Set ID or TYPEFULLY_SOCIAL_SET_ID",
    )
    parser.add_argument("--base-url", default="https://api.typefully.com", help=argparse.SUPPRESS)
    parser.add_argument("--mode", choices=("draft", "publish", "schedule"), default="draft")
    parser.add_argument("--publish-at", default=None, help="ISO 8601 time for --mode schedule")
    parser.add_argument(
        "--asset-root",
        "--image-root",
        dest="asset_root",
        default="output/assets",
        help="Root for numbered local media folders",
    )
    parser.add_argument("--db", default="output/state.sqlite3", help="SQLite checkpoint path; use --no-db to disable")
    parser.add_argument("--no-db", action="store_true", help="Do not update the SQLite checkpoint")
    parser.add_argument("--output-xlsx", default=None, help="Output workbook; default is <input>.typefully.xlsx")
    parser.add_argument("--in-place", action="store_true", help="Overwrite the input workbook")
    parser.add_argument("--limit", type=int, default=1, help="Rows to process in canary mode")
    parser.add_argument("--run-all", action="store_true", help="Process all eligible rows")
    parser.add_argument(
        "--row-number",
        action="append",
        type=int,
        default=None,
        help="Process these Excel row numbers only; may be repeated",
    )
    parser.add_argument("--execute", action="store_true", help="Perform real Typefully uploads; otherwise dry-run")
    parser.add_argument(
        "--list-social-sets",
        action="store_true",
        help="List accessible Typefully Social Sets and exit",
    )
    parser.add_argument(
        "--require-media",
        "--require-image",
        dest="require_media",
        action="store_true",
        help="Fail rows that do not have a local publishable video or fallback image",
    )
    parser.add_argument(
        "--only-status",
        default="SHARE_VALIDATED,ASSET_READY,EXPORTED",
        help="Comma-separated processing statuses eligible for Typefully",
    )
    return parser


def _output_path(input_path: Path, args: argparse.Namespace) -> Path:
    if args.in_place and args.output_xlsx:
        raise ValueError("--in-place and --output-xlsx cannot be used together")
    if args.in_place:
        return input_path
    if args.output_xlsx:
        return Path(args.output_xlsx)
    return input_path.with_name(f"{input_path.stem}.typefully{input_path.suffix}")


def main(argv: list[str] | None = None) -> int:
    _load_dotenv()
    args = _parser().parse_args(argv)
    if args.list_social_sets:
        if not args.api_key:
            raise ValueError("--list-social-sets requires TYPEFULLY_API_KEY in .env or the environment")
        response = TypefullyClient(args.api_key, base_url=args.base_url).list_social_sets()
        print(json.dumps(response, ensure_ascii=False, indent=2))
        return 0
    if not args.xlsx:
        raise ValueError("--xlsx is required unless --list-social-sets is used")
    if args.limit < 1:
        raise ValueError("--limit must be greater than zero")
    if args.row_number and any(row_number < 2 for row_number in args.row_number):
        raise ValueError("--row-number must refer to a data row (2 or greater)")
    if args.mode == "schedule" and not args.publish_at:
        raise ValueError("--mode schedule requires --publish-at")
    if args.execute and (not args.api_key or not args.social_set_id):
        raise ValueError("--execute requires TYPEFULLY_API_KEY and TYPEFULLY_SOCIAL_SET_ID")

    input_path = Path(args.xlsx)
    if not input_path.is_file():
        raise FileNotFoundError(f"Excel workbook not found: {input_path}")
    output_path = _output_path(input_path, args)
    workbook = load_workbook(input_path)
    sheet = workbook.active
    headers = _ensure_columns(sheet, TYPEFULLY_COLUMNS)
    status_column = next((headers[name] for name in STATUS_COLUMNS if name in headers), None)
    typefully_status_column = headers["Typefully状态"]
    only_status = {value.strip() for value in args.only_status.split(",") if value.strip()}

    candidates: list[int] = []
    requested_rows = set(args.row_number or ())
    for row_number in range(2, sheet.max_row + 1):
        if requested_rows and row_number not in requested_rows:
            continue
        row = _row_values(sheet, headers, row_number)
        if status_column is not None and str(sheet.cell(row=row_number, column=status_column).value or "") not in only_status:
            continue
        existing_status = str(sheet.cell(row=row_number, column=typefully_status_column).value or "")
        if existing_status in {"DRAFT_CREATED", "PUBLISH_REQUESTED", "PUBLISHED", "SCHEDULED"}:
            continue
        candidates.append(row_number)
    if not args.run_all and not requested_rows:
        candidates = candidates[: args.limit]

    print(f"typefully_candidates={len(candidates)} mode={args.mode} execute={int(args.execute)} output={output_path}")
    if not args.execute:
        for row_number in candidates:
            row = _row_values(sheet, headers, row_number)
            text = build_post_text(
                row,
                data_index=row_number - 1,
                asset_root=args.asset_root,
                workbook_path=input_path,
            )
            media = resolve_media_path(
                row,
                data_index=row_number - 1,
                workbook_path=input_path,
                asset_root=args.asset_root,
            )
            if args.require_media and media is None:
                sheet.cell(row=row_number, column=typefully_status_column, value="FAILED")
                sheet.cell(row=row_number, column=headers["Typefully错误信息"], value="未找到本地可发布视频或兜底图片")
                print(f"dry_run=row={row_number} status=FAILED media=-")
                continue
            media_error = validate_typefully_media(media)
            if media_error:
                sheet.cell(row=row_number, column=typefully_status_column, value="FAILED")
                sheet.cell(row=row_number, column=headers["Typefully错误信息"], value=media_error[:1000])
                print(f"dry_run=row={row_number} status=FAILED media={media} error={media_error}")
                continue
            sheet.cell(row=row_number, column=headers["X文案"], value=text)
            sheet.cell(row=row_number, column=headers["Typefully媒体路径"], value=str(media) if media else "")
            sheet.cell(row=row_number, column=typefully_status_column, value="DRY_RUN")
            sheet.cell(row=row_number, column=headers["Typefully错误信息"], value="")
            print(f"dry_run=row={row_number} media={media or '-'} text_length={len(text)}")
        workbook.save(output_path)
        return 0

    client = TypefullyClient(args.api_key, base_url=args.base_url)
    store = None if args.no_db else StateStore(args.db)
    for row_number in candidates:
        row = _row_values(sheet, headers, row_number)
        text = build_post_text(
            row,
            data_index=row_number - 1,
            asset_root=args.asset_root,
            workbook_path=input_path,
        )
        if not text:
            error = "X 文案为空"
            sheet.cell(row=row_number, column=typefully_status_column, value="FAILED")
            sheet.cell(row=row_number, column=headers["Typefully错误信息"], value=error)
            continue
        media = resolve_media_path(
            row,
            data_index=row_number - 1,
            workbook_path=input_path,
            asset_root=args.asset_root,
        )
        if args.require_media and media is None:
            error = "未找到本地可发布视频或兜底图片"
            sheet.cell(row=row_number, column=typefully_status_column, value="FAILED")
            sheet.cell(row=row_number, column=headers["Typefully错误信息"], value=error)
            continue
        media_error = validate_typefully_media(media)
        if media_error:
            sheet.cell(row=row_number, column=typefully_status_column, value="FAILED")
            sheet.cell(row=row_number, column=headers["Typefully错误信息"], value=media_error[:1000])
            sheet.cell(row=row_number, column=headers["Typefully更新时间"], value=_now())
            _record_state(store, row, status="TYPEFULLY_FAILED", metadata={}, error=media_error)
            print(f"row={row_number} status=FAILED error={media_error}")
            continue
        sheet.cell(row=row_number, column=headers["X文案"], value=text)
        sheet.cell(row=row_number, column=headers["Typefully媒体路径"], value=str(media) if media else "")
        try:
            publish_at = None
            if args.mode == "publish":
                publish_at = "now"
            elif args.mode == "schedule":
                publish_at = args.publish_at
            draft = client.create_x_draft(
                args.social_set_id,
                text=text,
                media_path=media,
                draft_title=_value(row, TITLE_COLUMNS) or None,
                publish_at=publish_at,
            )
            status = (
                "PUBLISH_REQUESTED"
                if args.mode == "publish"
                else "SCHEDULED"
                if args.mode == "schedule"
                else "DRAFT_CREATED"
            )
            sheet.cell(row=row_number, column=headers["Typefully草稿ID"], value=draft.draft_id)
            sheet.cell(row=row_number, column=headers["Typefully草稿链接"], value=draft.private_url or draft.share_url)
            sheet.cell(row=row_number, column=typefully_status_column, value=status)
            sheet.cell(row=row_number, column=headers["Typefully错误信息"], value="")
            sheet.cell(row=row_number, column=headers["Typefully更新时间"], value=_now())
            _record_state(
                store,
                row,
                status=f"TYPEFULLY_{status}",
                metadata={
                    "draft_id": draft.draft_id,
                    "private_url": draft.private_url,
                    "share_url": draft.share_url,
                    "published_url": draft.published_url,
                    "media_path": str(media) if media else "",
                },
            )
            print(f"row={row_number} status={status} draft_id={draft.draft_id}")
        except (TypefullyAPIError, OSError, ValueError) as exc:
            error = str(exc)
            sheet.cell(row=row_number, column=typefully_status_column, value="FAILED")
            sheet.cell(row=row_number, column=headers["Typefully错误信息"], value=error[:1000])
            sheet.cell(row=row_number, column=headers["Typefully更新时间"], value=_now())
            _record_state(store, row, status="TYPEFULLY_FAILED", metadata={}, error=error)
            print(f"row={row_number} status=FAILED error={error}")

    workbook.save(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
