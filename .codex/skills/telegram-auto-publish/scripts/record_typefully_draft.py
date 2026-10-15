#!/usr/bin/env python3
"""Record one MCP-verified Typefully draft without touching source columns."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--row-index", type=int, required=True)
    parser.add_argument("--row-number", type=int, required=True)
    parser.add_argument("--draft-id", type=int, required=True)
    parser.add_argument("--social-set-id", type=int, required=True)
    parser.add_argument("--media-path", type=Path, required=True)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[4]
    sys.path.insert(0, str(root / "src"))
    from xpan_auto.copy_review import validate_review
    from xpan_auto.publish_pipeline import FIXED_WORKBOOK_PATH, _headers, _row_identity, _row_values
    from xpan_auto.video_quality import inspect_publish_video

    asset_dir = root / "output" / "assets" / f"{args.row_index:04d}"
    media_input = args.media_path.absolute()
    media = media_input.resolve()
    if not media.is_file() or not 0 < media.stat().st_size < 50_000_000:
        parser.error("final media is missing or exceeds the Bot API size gate")

    workbook = load_workbook(FIXED_WORKBOOK_PATH)
    temporary = None
    try:
        sheet = workbook.active
        headers = _headers(sheet)
        if not 2 <= args.row_number <= sheet.max_row:
            parser.error("row number outside workbook")
        row = _row_values(sheet, headers, args.row_number)
        serial, title, description, link = _row_identity(row, args.row_number)
        if serial != args.row_index:
            parser.error("serial does not match Excel row")
        expected_media = Path(str(row.get("媒体路径") or "")).absolute()
        if media_input != expected_media or media.parent != asset_dir.resolve():
            parser.error("media path does not match the selected row")
        media_type = str(row.get("媒体类型") or "")
        if media_type == "video":
            if not media_input.name.startswith("project-video-original."):
                parser.error("video is not the final prepared media")
            qa = inspect_publish_video(media, max_duration_seconds=600, max_size_bytes=49_999_999)
            if not qa.passed:
                parser.error("final video QA failed")
        elif media_type == "image":
            if media_input.name != "cover-v1-information-card.png":
                parser.error("image is not the final prepared cover")
            quality = json.loads((asset_dir / "telegram-quality.json").read_text(encoding="utf-8"))
            if not quality.get("passed"):
                parser.error("final cover QA failed")
            with Image.open(media) as image:
                image.verify()
        else:
            parser.error("selected row has no supported media type")
        ledger = json.loads((asset_dir / "typefully-upload-state.json").read_text(encoding="utf-8"))
        if (ledger.get("serial") != serial or ledger.get("excel_row_number") != args.row_number
                or ledger.get("social_set_id") != args.social_set_id
                or ledger.get("media_file") != media_input.name or not ledger.get("media_id")
                or hashlib.sha256(media.read_bytes()).hexdigest() != ledger.get("media_sha256")):
            parser.error("Typefully upload state does not match the reviewed final media")
        validate_review(root / "output" / "assets" / f"{serial:04d}", serial=serial,
                        title=title, description=description, link=link)
        existing = row.get("Typefully草稿ID")
        if existing and str(existing) != str(args.draft_id):
            parser.error("another Typefully draft ID is already recorded")
        if not row.get("Telegram消息ID") or str(row.get("Telegram发布状态") or "") not in {"已发送", "SENT", "TELEGRAM_SENT"}:
            parser.error("Telegram success is not recorded")
        required = ("Typefully媒体路径", "Typefully草稿ID", "Typefully草稿链接",
                    "Typefully状态", "Typefully错误信息", "Typefully更新时间",
                    "X发布状态", "X错误信息", "统一发布状态", "统一发布错误信息",
                    "统一发布更新时间")
        if any(name not in headers for name in required):
            parser.error("required output column is missing")
        now = datetime.now(timezone.utc).isoformat()
        values = {
            "Typefully媒体路径": str(media_input),
            "Typefully草稿ID": str(args.draft_id),
            "Typefully草稿链接": f"https://typefully.com/?d={args.draft_id}&a={args.social_set_id}",
            "Typefully状态": "DRAFT_CREATED",
            "Typefully错误信息": None,
            "Typefully更新时间": now,
            "X发布状态": "待发布",
            "X错误信息": None,
            "统一发布状态": "COMPLETED",
            "统一发布错误信息": None,
            "统一发布更新时间": now,
        }
        for name, value in values.items():
            sheet.cell(args.row_number, headers[name]).value = value
        with tempfile.NamedTemporaryFile(suffix=".xlsx", dir=FIXED_WORKBOOK_PATH.parent, delete=False) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        temporary.replace(FIXED_WORKBOOK_PATH)
        print(f"serial={serial} row={args.row_number} draft_id={args.draft_id} status=DRAFT_CREATED")
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
