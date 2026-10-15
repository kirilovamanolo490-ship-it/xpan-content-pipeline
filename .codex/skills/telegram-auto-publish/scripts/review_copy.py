#!/usr/bin/env python3
"""Record or check source-backed editorial approval for one prepared project."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def project_root() -> Path:
    return Path(__file__).resolve().parents[4]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("action", choices=("record", "check"))
    parser.add_argument("--row-index", type=int, required=True, help="Project serial, not Excel row number")
    parser.add_argument("--row-number", type=int, required=True, help="Actual 1-based Excel row number")
    parser.add_argument("--category", choices=("commercial", "tool", "course", "reference", "media", "other"))
    parser.add_argument("--style-label")
    parser.add_argument("--angle")
    parser.add_argument("--key-phrase", help="One factual phrase that appears in both channel copies")
    parser.add_argument("--source-note", help="Concrete finding from the workbook description, full tree or preview")
    parser.add_argument("--evidence", action="append", default=[], help="Exact name or path from project details; repeat")
    parser.add_argument("--offer", default="", help="Commercial item: what is produced or delivered")
    parser.add_argument("--audience", default="", help="Commercial item: who would use or buy it")
    parser.add_argument("--revenue", default="", help="Commercial item: income source or explicitly unverified link")
    parser.add_argument("--support", choices=("supported", "partial", "unclear"))
    args = parser.parse_args()
    if args.row_index < 1 or args.row_number < 2:
        parser.error("--row-index must be positive and --row-number must be at least 2")
    sys.path.insert(0, str(project_root() / "src"))
    from xpan_auto.copy_review import record_review, validate_review
    from xpan_auto.publish_pipeline import FIXED_WORKBOOK_PATH, _headers, _row_identity, _row_values

    workbook = load_workbook(FIXED_WORKBOOK_PATH, read_only=True)
    try:
        sheet = workbook.active
        if args.row_number > sheet.max_row:
            parser.error("--row-number is outside the fixed workbook")
        serial, title, description, link = _row_identity(
            _row_values(sheet, _headers(sheet), args.row_number), args.row_number
        )
    finally:
        workbook.close()
    if serial != args.row_index:
        parser.error("--row-index does not match the serial at --row-number")

    asset_dir = project_root() / "output" / "assets" / f"{args.row_index:04d}"
    if args.action == "record":
        for flag in ("category", "style_label", "angle", "key_phrase", "source_note"):
            if not getattr(args, flag):
                parser.error(f"record requires --{flag.replace('_', '-')}")
        result = record_review(
            asset_dir,
            category=args.category,
            style_label=args.style_label,
            primary_angle=args.angle,
            key_phrase=args.key_phrase,
            description=description,
            source_note=args.source_note,
            source_evidence=args.evidence,
            offer=args.offer,
            audience=args.audience,
            revenue=args.revenue,
            support=args.support or "",
        )
    else:
        result = validate_review(
            asset_dir,
            serial=args.row_index,
            title=title,
            description=description,
            link=link,
        )
    print(json.dumps({"review": "passed", **result}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
