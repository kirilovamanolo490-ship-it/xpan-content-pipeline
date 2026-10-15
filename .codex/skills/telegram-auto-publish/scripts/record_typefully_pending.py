#!/usr/bin/env python3
"""Record an unresolved MCP media upload without losing its retry identity."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--row-index", type=int, required=True)
    parser.add_argument("--row-number", type=int, required=True)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[4]
    sys.path.insert(0, str(root / "src"))
    from xpan_auto.publish_pipeline import FIXED_WORKBOOK_PATH, _headers, _row_identity, _row_values

    asset_dir = root / "output" / "assets" / f"{args.row_index:04d}"
    ledger = json.loads((asset_dir / "typefully-upload-state.json").read_text(encoding="utf-8"))
    media = asset_dir / str(ledger.get("media_file"))
    if (ledger.get("serial") != args.row_index or ledger.get("excel_row_number") != args.row_number
            or not ledger.get("media_id") or ledger.get("last_verified_status") != "processing"
            or hashlib.sha256(media.read_bytes()).hexdigest() != ledger.get("media_sha256")):
        parser.error("pending upload ledger does not match the row or media")

    workbook = load_workbook(FIXED_WORKBOOK_PATH)
    temporary = None
    try:
        sheet = workbook.active
        headers = _headers(sheet)
        row = _row_values(sheet, headers, args.row_number)
        serial, _title, _description, _link = _row_identity(row, args.row_number)
        if serial != args.row_index or row.get("Typefully草稿ID"):
            parser.error("row identity changed or draft already exists")
        if not row.get("Telegram消息ID"):
            parser.error("Telegram message ID is missing")
        now = datetime.now(timezone.utc).isoformat()
        reason = f"Typefully 媒体 {ledger['media_id']} 仍在 processing；按同一 ID 回读后再建稿"
        values = {
            "Typefully媒体路径": str(media.resolve()),
            "Typefully状态": "RESULT_UNKNOWN",
            "Typefully错误信息": reason,
            "Typefully更新时间": now,
            "X发布状态": "结果待确认",
            "X错误信息": reason,
            "统一发布状态": "TELEGRAM_SENT_X_PENDING",
            "统一发布错误信息": reason,
            "统一发布更新时间": now,
        }
        if any(name not in headers for name in values):
            parser.error("required output column is missing")
        for name, value in values.items():
            sheet.cell(args.row_number, headers[name], value)
        with tempfile.NamedTemporaryFile(suffix=".xlsx", dir=FIXED_WORKBOOK_PATH.parent, delete=False) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        temporary.replace(FIXED_WORKBOOK_PATH)
        print(f"serial={serial} row={args.row_number} media_id={ledger['media_id']} status=RESULT_UNKNOWN")
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
