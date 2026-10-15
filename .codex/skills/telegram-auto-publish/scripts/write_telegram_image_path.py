#!/usr/bin/env python3
"""Compatibility CLI for writing one verified image path without COM."""
import argparse
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[4] / "src"))
from openpyxl import load_workbook
from xpan_auto.publish_pipeline import FIXED_WORKBOOK_PATH, _headers, _write_image_path
from xpan_auto.runtime import resolve_project_path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--row-index", type=int, required=True, help="Project serial, not worksheet row")
    parser.add_argument("--image-path", required=True)
    parser.add_argument("--workbook-path", default=str(FIXED_WORKBOOK_PATH))
    args = parser.parse_args()
    workbook_path = resolve_project_path(args.workbook_path).resolve()
    if workbook_path != FIXED_WORKBOOK_PATH:
        parser.error("Only the fixed workflow workbook is allowed")
    wb = load_workbook(workbook_path)
    try:
        sheet = wb["处理结果"]; headers = _headers(sheet)
        rows = [row for row in range(2, sheet.max_row + 1)
                if str(sheet.cell(row, headers["序号"]).value) == str(args.row_index)]
    finally:
        wb.close()
    if len(rows) != 1:
        parser.error("Serial is missing or ambiguous")
    _write_image_path(workbook_path, rows[0], args.row_index, resolve_project_path(args.image_path))
    print(f"image_path_updated serial={args.row_index} row={rows[0]}")


if __name__ == "__main__":
    main()
