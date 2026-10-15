import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from xpan_auto.runtime import resolve_project_path
from xpan_auto.publish_pipeline import _write_image_path, PublishPipelineError
from openpyxl import Workbook, load_workbook


class RuntimeTests(unittest.TestCase):
    def test_moved_windows_paths_keep_full_project_relative_identity(self):
        with TemporaryDirectory() as directory:
            root = Path(directory) / "X网盘拉新"
            expected = root / "output/assets/0170/video.mp4"
            self.assertEqual(resolve_project_path(r"F:\X网盘拉新\output\assets\0170\video.mp4", root=root), expected)
            self.assertEqual(resolve_project_path(r"output\assets\0170\video.mp4", root=root), expected)
            self.assertEqual(resolve_project_path("/Users/old/X网盘拉新/output/assets/0170/video.mp4", root=root), expected)

    def test_unrelated_drive_path_is_not_replaced_by_same_named_file(self):
        with TemporaryDirectory() as directory:
            root = Path(directory) / "project"
            value = r"C:\another-project\output\video.mp4"
            self.assertEqual(resolve_project_path(value, root=root), Path(value))

    def test_cache_relative_path_and_project_relative_path(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            base = root / "output/assets/0001"
            self.assertEqual(resolve_project_path("video.mp4", base=base, root=root), base / "video.mp4")
            self.assertEqual(resolve_project_path("output/assets/0001/video.mp4", base=base, root=root), base / "video.mp4")

    def test_image_write_uses_actual_row_and_preserves_other_cells(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "book.xlsx"
            image = Path(directory) / "cover.png"
            image.write_bytes(b"image")
            wb = Workbook(); ws = wb.active; ws.title = "处理结果"
            ws.append(["序号", "Telegram消息ID", "图片", "Typefully状态"])
            ws.append([10, "123", "old", "DRAFT_CREATED"])
            ws.append([42, "456", "other", "DRAFT_CREATED"])
            ws["C3"].number_format = "@"
            wb.save(path); wb.close()
            _write_image_path(path, 3, 42, image)
            wb = load_workbook(path); ws = wb.active
            self.assertEqual(ws["C3"].value, str(image.resolve()))
            self.assertEqual(ws["C3"].number_format, "@")
            self.assertEqual(ws["C2"].value, "old")
            self.assertEqual(ws["B3"].value, "456")
            self.assertEqual(ws["D3"].value, "DRAFT_CREATED")
            wb.close()
            before = path.read_bytes()
            with self.assertRaises(PublishPipelineError):
                _write_image_path(path, 2, 42, image)
            self.assertEqual(path.read_bytes(), before)
