import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from xpan_auto.excel_exporter import export_rows, update_existing_row


class ExcelExporterTests(unittest.TestCase):
    def test_export_creates_expected_headers(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output = export_rows(
                [{"序号": 1, "原夸克链接": "https://pan.quark.cn/s/abc123", "处理状态": "EXTRACTED"}],
                Path(temp_dir) / "result.xlsx",
            )

            workbook = load_workbook(output)
            sheet = workbook["处理结果"]
            headers = [cell.value for cell in sheet[1]]

            self.assertEqual(headers[0], "序号")
            self.assertIn("新夸克链接", headers)
            self.assertEqual(sheet["D2"].value, "https://pan.quark.cn/s/abc123")

    def test_update_existing_row_preserves_publisher_columns(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output = export_rows(
                [{"序号": 1, "原夸克链接": "https://pan.quark.cn/s/old", "处理状态": "EXTRACTED"}],
                Path(temp_dir) / "result.xlsx",
            )
            workbook = load_workbook(output)
            sheet = workbook.active
            sheet.cell(row=1, column=17, value="Typefully状态")
            sheet.cell(row=2, column=17, value="DRAFT_CREATED")
            workbook.save(output)

            update_existing_row(
                2,
                {
                    "序号": 1,
                    "原夸克链接": "https://pan.quark.cn/s/old",
                    "新夸克链接": "https://pan.quark.cn/s/new",
                    "处理状态": "SHARE_VALIDATED",
                },
                output,
            )

            updated = load_workbook(output)
            self.assertEqual(updated.active["F2"].value, "https://pan.quark.cn/s/new")
            self.assertEqual(updated.active.cell(row=2, column=17).value, "DRAFT_CREATED")
