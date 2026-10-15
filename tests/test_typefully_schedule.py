import hashlib
import json
import os
import stat
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill

from xpan_auto import typefully_schedule as schedule


class SchedulePlanTests(unittest.TestCase):
    def test_malformed_items_are_not_silently_treated_as_empty_queue(self):
        bad_snapshots = [
            [{}],
            {"days": [{"items": [{"kind": "custom_time", "draft": {}}]}]},
            {"results": [{"id": 12345}]},
        ]
        for snapshot in bad_snapshots:
            with self.subTest(snapshot=snapshot), self.assertRaises(ValueError):
                schedule.build_plan("2026-10-13", "2026-10-13T17:00:00+08:00", snapshot)

    def plan(self, now, queue=None, drafts=None):
        snapshot = {"queue": queue or {"days": []}, "list": {"results": drafts or []}}
        return schedule.build_plan("2026-10-13", now, snapshot)

    def test_fixed_eight_slots_keep_2000_empty(self):
        plan = self.plan("2026-10-12T18:00:00+08:00")
        self.assertEqual([slot["time"] for slot in plan["slots"]], [
            "19:30", "19:45", "20:15", "20:30", "20:45", "21:00", "21:15", "21:30",
        ])
        self.assertNotIn("20:00", json.dumps(plan))
        self.assertEqual([slot["status"] for slot in plan["slots"]], ["available"] * 8)

    def test_after_eight_pm_only_future_fixed_slots_are_available(self):
        plan = self.plan("2026-10-13T20:05:00+08:00")
        available = [slot["time"] for slot in plan["slots"] if slot["status"] == "available"]
        self.assertEqual(available, ["20:15", "20:30", "20:45", "21:00", "21:15", "21:30"])
        self.assertEqual([slot["status"] for slot in plan["slots"][:2]], ["expired", "expired"])

    def test_scheduled_and_published_occupy_slots_but_planned_is_labeled_only(self):
        drafts = [
            {"id": "a", "status": "scheduled", "scheduled_date": "2026-10-13T19:30:00+08:00"},
            {"id": "b", "status": "published", "published_at": "2026-10-13T12:15:00Z"},
            {"id": "c", "status": "planned", "scheduled_date": "2026-10-13T20:30:00+08:00"},
        ]
        plan = self.plan("2026-10-13T18:00:00+08:00", drafts=drafts)
        statuses = {slot["time"]: slot["status"] for slot in plan["slots"]}
        self.assertEqual(statuses["19:30"], "occupied_scheduled")
        self.assertEqual(statuses["20:15"], "occupied_published")
        self.assertEqual(statuses["20:30"], "occupied_planned")
        self.assertEqual(len([slot for slot in plan["slots"] if slot["status"].startswith("occupied")]), 3)

    def test_cross_timezone_slot_conversion(self):
        drafts = [{"id": "a", "status": "scheduled", "scheduled_date": "2026-10-13T11:30:00Z"}]
        plan = self.plan("2026-10-13T18:00:00+08:00", drafts=drafts)
        self.assertEqual(next(s for s in plan["slots"] if s["time"] == "19:30")["status"], "occupied_scheduled")

    def test_accepts_structured_queue_items_and_deduplicates_same_draft(self):
        draft = {"id": "a", "status": "scheduled", "scheduled_date": "2026-10-13T19:30:00+08:00"}
        snapshot = {
            "queue": {"structuredContent": {"days": [{"items": [{"draft": draft}]}]}},
            "list": {"structuredContent": {"results": [draft]}},
        }
        plan = schedule.build_plan("2026-10-13", "2026-10-13T18:00:00+08:00", snapshot)
        occupied = [slot for slot in plan["slots"] if slot["status"].startswith("occupied")]
        self.assertEqual(len(occupied), 1)
        self.assertEqual(occupied[0]["draft_id"], "a")
        direct = schedule.build_plan("2026-10-13", "2026-10-13T18:00:00+08:00", [draft])
        self.assertEqual(next(slot for slot in direct["slots"] if slot["time"] == "19:30")["status"], "occupied_scheduled")

    def test_ignores_real_empty_queue_slots_and_rejects_malformed_queue_objects(self):
        snapshot = {"days": [{"items": [
            {"kind": "queue_slot", "at": "2026-10-13T19:30:00+08:00", "draft": None},
        ]}]}
        plan = schedule.build_plan("2026-10-13", "2026-10-13T18:00:00+08:00", snapshot)
        self.assertEqual(next(s for s in plan["slots"] if s["time"] == "19:30")["status"], "available")
        malformed = {"days": [{"items": [{"kind": "queue_slot", "at": "...", "draft": "bad"}]}]}
        with self.assertRaisesRegex(ValueError, "draft"):
            schedule.build_plan("2026-10-13", "2026-10-13T18:00:00+08:00", malformed)

    def test_accepts_documented_standalone_results_shape(self):
        result = schedule.build_plan("2026-10-13", "2026-10-13T18:00:00+08:00", {"results": []})
        self.assertEqual(len(result["slots"]), 8)

    def test_rejects_paginated_or_truncated_snapshots(self):
        snapshots = [
            {"days": [], "next": "cursor-2"},
            {"days": [], "truncated": True},
            {"results": [], "next": "cursor-2"},
            {"queue": {"days": []}, "list": {"results": [], "truncated": True}},
        ]
        for snapshot in snapshots:
            with self.subTest(snapshot=snapshot), self.assertRaisesRegex(ValueError, "完整|分页|截断"):
                schedule.build_plan("2026-10-13", "2026-10-13T18:00:00+08:00", snapshot)

    def test_rejects_conflicting_duplicate_ids_and_incomplete_snapshot(self):
        duplicate = [
            {"id": "same", "status": "scheduled", "scheduled_date": "2026-10-13T19:30:00+08:00"},
            {"id": "same", "status": "scheduled", "scheduled_date": "2026-10-13T19:45:00+08:00"},
        ]
        with self.assertRaisesRegex(ValueError, "冲突"):
            self.plan("2026-10-13T18:00:00+08:00", drafts=duplicate)
        with self.assertRaisesRegex(ValueError, "完整"):
            schedule.build_plan("2026-10-13", "2026-10-13T18:00:00+08:00", {"queue": {"error": "timeout"}})


class ScheduleRecordTests(unittest.TestCase):
    def setUp(self):
        env_patch = patch.dict(os.environ, {"XPAN_TYPEFULLY_SOCIAL_SET_ID": "123456"})
        env_patch.start()
        self.addCleanup(env_patch.stop)
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.workbook_path = self.root / "output" / "quark_links.typefully.xlsx"
        self.workbook_path.parent.mkdir()
        self.asset_dir = self.workbook_path.parent / "assets" / "0007"
        self.asset_dir.mkdir(parents=True)
        self.link = "https://pan.quark.cn/s/example"
        self.copy = "这里是经过审阅的文案，介绍一个清楚的项目关键点。\n\n" + self.link
        (self.asset_dir / "x.txt").write_text(self.copy, encoding="utf-8")
        (self.asset_dir / "telegram-project-details.json").write_text(json.dumps({
            "row_index": 7, "title": "示例项目", "source_share_link": self.link,
            "summary": {"truncated": False}, "items": [{"type": "file", "name": "示例项目内容.pdf", "path": "示例项目内容.pdf"}],
        }, ensure_ascii=False), encoding="utf-8")
        (self.asset_dir / "telegram-caption.html").write_text(
            "示例项目\n\n介绍一个清楚的项目关键点。\n\n🔗 " + self.link +
            "\n\n#示例 #项目\n\n<a href=\"https://example.invalid/x/profile\">X</a> <a href=\"https://example.invalid/telegram/channel\">频道</a>",
            encoding="utf-8")
        from xpan_auto.copy_review import record_review
        record_review(self.asset_dir, category="reference", style_label="疑问／观点型",
                      primary_angle="解释示例项目的实际内容", key_phrase="项目关键点",
                      description="介绍示例项目内容。", source_note="完整目录包含示例项目内容文件。",
                      source_evidence=["示例项目内容.pdf"])
        media = b"stable prepared media"
        (self.asset_dir / "prepared.mp4").write_bytes(media)
        (self.asset_dir / "typefully-upload-state.json").write_text(json.dumps({
            "serial": 7, "excel_row_number": 2, "social_set_id": 123456,
            "media_file": "prepared.mp4", "media_id": "media-7",
            "media_sha256": hashlib.sha256(media).hexdigest(),
        }), encoding="utf-8")
        self.headers = ["序号", "资料标题", "文字介绍", "新夸克链接", "Typefully草稿ID",
                        "Typefully状态", "Typefully更新时间", "X发布状态", "X排期时间",
                        "X发布时间", "X帖子链接", "X错误信息", "Typefully错误信息",
                        "Telegram发布状态", "Telegram消息ID", "备注"]
        book = Workbook()
        sheet = book.active
        sheet.title = "处理结果"
        sheet.append(self.headers)
        sheet.append([7, "示例项目", "介绍示例项目内容。", self.link, "1010101",
                      "DRAFT_CREATED", "old-update", "待发布", "", "", "", "", "",
                      "已发送", "tg-77", "保留此数据"])
        sheet.cell(2, len(self.headers))._style = sheet.cell(1, len(self.headers))._style
        sheet["P2"].hyperlink = "https://example.com/keep"
        sheet["P2"].comment = Comment("保留批注", "Codex")
        sheet.conditional_formatting.add("B2:B2", CellIsRule(operator="equal", formula=['"示例项目"'], fill=PatternFill("solid", fgColor="FFFF00")))
        sheet.merge_cells("R1:S1")
        sheet.row_dimensions[2].height = 27
        sheet.column_dimensions["P"].width = 22
        book.save(self.workbook_path)
        book.close()
        with ZipFile(self.workbook_path, "a") as archive:
            archive.writestr("xl/media/preserved-marker.bin", b"media must remain byte-for-byte")
            archive.comment = b"keep archive comment"
        os.chmod(self.workbook_path, 0o640)
        with ZipFile(self.workbook_path) as archive:
            self.package_before = {name: archive.read(name) for name in archive.namelist()}

    def tearDown(self):
        self.temp.cleanup()

    def draft(self, **changes):
        data = {
            "id": 1010101, "social_set_id": 123456, "status": "scheduled",
            "scheduled_date": "2026-10-13T19:30:00+08:00", "updated_at": "2026-10-12T10:11:12Z",
            "platforms": self.platforms(),
        }
        data.update(changes)
        return {"structuredContent": data}

    def platforms(self, *, text=None, media_ids=None, x_enabled=True, linkedin_enabled=False):
        return {"x": {"enabled": x_enabled, "posts": [{
                    "text": self.copy if text is None else text,
                    "media_ids": ["media-7"] if media_ids is None else media_ids,
                }]},
                "linkedin": {"enabled": linkedin_enabled, "posts": []}, "x_article": None}

    def record(self, draft=None, expected="2026-10-13T19:30:00+08:00"):
        with patch.object(schedule, "FIXED_WORKBOOK_PATH", self.workbook_path), \
             patch.object(schedule, "PROJECT_ROOT", self.root):
            return schedule.record_schedule(row_index=7, row_number=2, draft_json=draft or self.draft(),
                                            expected_time=expected, social_set_id=123456)

    def values(self):
        book = load_workbook(self.workbook_path)
        try:
            sheet = book.active
            return {self.headers[i]: sheet.cell(2, i + 1).value for i in range(len(self.headers))}
        finally:
            book.close()

    def test_rejects_bad_time_id_text_and_media_without_touching_master(self):
        before = self.workbook_path.read_bytes()
        for draft, expected in [
            (self.draft(), "2026-10-13T19:45:00+08:00"),
            (self.draft(id="wrong"), "2026-10-13T19:30:00+08:00"),
            (self.draft(social_set_id=123), "2026-10-13T19:30:00+08:00"),
            (self.draft(platforms=self.platforms(text="different")), "2026-10-13T19:30:00+08:00"),
            (self.draft(platforms=self.platforms(media_ids=["other"])), "2026-10-13T19:30:00+08:00"),
            (self.draft(platforms=self.platforms(x_enabled=False)), "2026-10-13T19:30:00+08:00"),
            (self.draft(platforms=self.platforms(linkedin_enabled=True)), "2026-10-13T19:30:00+08:00"),
        ]:
            with self.assertRaises(ValueError):
                self.record(draft, expected)
            self.assertEqual(self.workbook_path.read_bytes(), before)

    def test_rejects_any_real_published_evidence(self):
        before = self.workbook_path.read_bytes()
        for field, value in [("published_at", "2026-10-13T19:30:00+08:00"),
                             ("x_post_published_at", "2026-10-13T19:30:00+08:00"),
                             ("x_published_url", "https://x.com/post/123")]:
            with self.subTest(field=field), self.assertRaisesRegex(ValueError, "发布|published"):
                self.record(self.draft(**{field: value}))
            self.assertEqual(self.workbook_path.read_bytes(), before)

    def test_rejects_existing_post_url_in_workbook(self):
        book = load_workbook(self.workbook_path)
        book["处理结果"]["K2"] = "https://x.com/post/already-published"
        book.save(self.workbook_path)
        book.close()
        before = self.workbook_path.read_bytes()
        with self.assertRaisesRegex(ValueError, "发布结果"):
            self.record()
        self.assertEqual(self.workbook_path.read_bytes(), before)

    def test_records_when_schedule_cell_is_absent_before_later_populated_columns(self):
        book = load_workbook(self.workbook_path)
        sheet = book.active
        sheet.cell(2, self.headers.index("X排期时间") + 1).value = None
        book.save(self.workbook_path)
        book.close()
        result = self.record()
        self.assertEqual(result["status"], "SCHEDULED")
        self.assertEqual(self.values()["X排期时间"], "2026-10-13T19:30:00+08:00")
        self.assertEqual(self.values()["备注"], "保留此数据")

    def test_updates_only_schedule_metadata_and_preserves_styles_and_telegram(self):
        before = load_workbook(self.workbook_path)
        before_sheet = before.active
        styles = {(r, c): tuple(before_sheet.cell(r, c)._style) if before_sheet.cell(r, c).has_style else (0,) * 9
                  for r in range(1, 3) for c in range(1, len(self.headers) + 1)}
        values = {(r, c): before_sheet.cell(r, c).value for r in range(1, 3) for c in range(1, len(self.headers) + 1)}
        before.close()
        result = self.record()
        self.assertEqual(result["status"], "SCHEDULED")
        after = load_workbook(self.workbook_path)
        try:
            sheet = after.active
            h = {cell.value: idx for idx, cell in enumerate(sheet[1], 1)}
            self.assertEqual(sheet.cell(2, h["Typefully状态"]).value, "SCHEDULED")
            self.assertEqual(sheet.cell(2, h["Typefully更新时间"]).value, "2026-10-12T10:11:12Z")
            self.assertEqual(sheet.cell(2, h["X发布状态"]).value, "已排期")
            self.assertEqual(sheet.cell(2, h["X排期时间"]).value, "2026-10-13T19:30:00+08:00")
            for name in ("X发布时间", "X帖子链接"):
                self.assertIn(sheet.cell(2, h[name]).value, (None, ""))
            for name in ("Telegram发布状态", "Telegram消息ID", "备注"):
                self.assertEqual(sheet.cell(2, h[name]).value, values[(2, h[name])])
            for r in range(1, 3):
                for c in range(1, len(self.headers) + 1):
                    signature = tuple(sheet.cell(r, c)._style) if sheet.cell(r, c).has_style else (0,) * 9
                    self.assertEqual(signature, styles[(r, c)])
                    if (r, c) not in {(2, h["Typefully状态"]), (2, h["Typefully更新时间"]),
                                      (2, h["X发布状态"]), (2, h["X排期时间"]),
                                      (2, h["X错误信息"]), (2, h["Typefully错误信息"])}:
                        self.assertEqual(sheet.cell(r, c).value, values[(r, c)])
        finally:
            after.close()
        with ZipFile(self.workbook_path) as archive:
            self.assertTrue(archive.testzip() is None)
            self.assertEqual(archive.comment, b"keep archive comment")
            self.assertEqual(archive.read("xl/media/preserved-marker.bin"), b"media must remain byte-for-byte")
            package_after = {name: archive.read(name) for name in archive.namelist()}
        self.assertEqual(set(package_after), set(self.package_before))
        changed_parts = [name for name in package_after if package_after[name] != self.package_before[name]]
        self.assertEqual(len(changed_parts), 1)
        self.assertTrue(changed_parts[0].startswith("xl/worksheets/"))
        self.assertEqual(stat.S_IMODE(self.workbook_path.stat().st_mode), 0o640)
        check = load_workbook(self.workbook_path)
        try:
            sheet = check["处理结果"]
            self.assertEqual(sheet["P2"].hyperlink.target, "https://example.com/keep")
            self.assertEqual(sheet["P2"].comment.text, "保留批注")
            self.assertEqual({str(item) for item in sheet.merged_cells.ranges}, {"R1:S1"})
            self.assertEqual(sheet.row_dimensions[2].height, 27)
            self.assertEqual(sheet.column_dimensions["P"].width, 22)
            self.assertEqual(len(sheet.conditional_formatting), 1)
        finally:
            check.close()

    def test_same_draft_same_slot_is_idempotent(self):
        self.record()
        first = self.workbook_path.read_bytes()
        second = self.record()
        self.assertTrue(second["idempotent"])
        self.assertEqual(self.workbook_path.read_bytes(), first)

    def test_changed_or_missing_media_cannot_be_recorded(self):
        before = self.workbook_path.read_bytes()
        media = self.asset_dir / "prepared.mp4"
        media.write_bytes(b"changed after review")
        with self.assertRaisesRegex(ValueError, "SHA-256"):
            self.record()
        self.assertEqual(self.workbook_path.read_bytes(), before)
        media.unlink()
        with self.assertRaisesRegex(ValueError, "媒体"):
            self.record()
        self.assertEqual(self.workbook_path.read_bytes(), before)

    def test_missing_hash_cannot_be_recorded(self):
        before = self.workbook_path.read_bytes()
        state_path = self.asset_dir / "typefully-upload-state.json"
        state = json.loads(state_path.read_text())
        state.pop("media_sha256")
        state_path.write_text(json.dumps(state))
        with self.assertRaisesRegex(ValueError, "SHA-256"):
            self.record()
        self.assertEqual(self.workbook_path.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
