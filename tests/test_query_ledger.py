import hashlib
import tempfile
import unittest
import sys
from argparse import Namespace
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import query_ledger


class QueryLedgerTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "ledger.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "处理结果"
        sheet.append(["序号", "资料标题", "原夸克链接", "处理状态"])
        sheet.append([21, "标题甲", "https://pan.quark.cn/s/a", "完成"])
        sheet.append([22, "标题乙", "https://pan.quark.cn/s/b", "待处理"])
        sheet.append([None, None, None, None])
        sheet.append([23, "标题丙", "https://pan.quark.cn/s/c", "完成"])
        workbook.save(self.path)
        workbook.close()
        self.patch = patch.object(query_ledger, "WORKBOOK_PATH", self.path)
        self.patch.start()

    def tearDown(self):
        self.patch.stop()
        self.tmp.cleanup()

    @staticmethod
    def args(**kwargs):
        values = dict(sequence=None, row_number=None, columns=None,
                      contains=None, where=[], limit=None, summary=False)
        values.update(kwargs)
        return Namespace(**values)

    def test_sequence_range_returns_sequence_and_real_excel_row(self):
        result = query_ledger.query_ledger(self.args(sequence="22-23"))
        self.assertEqual([row["sequence"] for row in result["rows"]], [22, 23])
        self.assertEqual([row["excel_row"] for row in result["rows"]], [3, 5])
        self.assertEqual(result["summary"]["data_rows"], 3)

    def test_exact_filter_keyword_columns_and_limit_combine(self):
        result = query_ledger.query_ledger(self.args(
            columns="资料标题", contains="标题", where=["处理状态=完成"], limit=1
        ))
        self.assertEqual(result["summary"]["matched_rows"], 2)
        self.assertEqual(len(result["rows"]), 1)
        self.assertEqual(result["rows"][0]["data"]["资料标题"], "标题甲")
        self.assertEqual(result["rows"][0]["excel_row"], 2)

    def test_real_row_number_selector(self):
        result = query_ledger.query_ledger(self.args(row_number=5))
        self.assertEqual(len(result["rows"]), 1)
        self.assertEqual(result["rows"][0]["sequence"], 23)

    def test_unknown_column_fails_clearly(self):
        with self.assertRaisesRegex(ValueError, "未知列：不存在字段"):
            query_ledger.query_ledger(self.args(columns="不存在字段"))

    def test_no_arguments_returns_summary_only_and_counts_blank_rows(self):
        result = query_ledger.query_ledger(self.args())
        self.assertEqual(result["rows"], [])
        self.assertEqual(result["summary"]["data_rows"], 3)
        self.assertEqual(result["summary"]["returned_rows"], 0)
        self.assertFalse(result["summary"]["truncated"])

    def test_filtered_query_defaults_to_twenty_rows_and_reports_truncation(self):
        workbook = load_workbook(self.path)
        sheet = workbook["处理结果"]
        for number in range(24, 46):
            sheet.append([number, f"标题{number}", "link", "完成"])
        workbook.save(self.path)
        workbook.close()
        result = query_ledger.query_ledger(self.args(contains="标题"))
        self.assertEqual(len(result["rows"]), 20)
        self.assertEqual(result["summary"]["matched_rows"], 25)
        self.assertEqual(result["summary"]["returned_rows"], 20)
        self.assertTrue(result["summary"]["truncated"])

    def test_invalid_sequence_range_and_filter_field_fail_clearly(self):
        with self.assertRaisesRegex(ValueError, "无效序号范围"):
            query_ledger.query_ledger(self.args(sequence="5-2"))
        with self.assertRaisesRegex(ValueError, "未知字段：不存在字段"):
            query_ledger.query_ledger(self.args(where=["不存在字段=x"]))

    def test_query_does_not_modify_workbook(self):
        before = hashlib.sha256(self.path.read_bytes()).hexdigest()
        query_ledger.query_ledger(self.args(sequence="21"))
        after = hashlib.sha256(self.path.read_bytes()).hexdigest()
        self.assertEqual(before, after)


if __name__ == "__main__":
    unittest.main()
