import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from xpan_auto.publish_pipeline import (
    FIXED_WORKBOOK_PATH,
    PUBLISH_MEDIA_POLICY_VERSION,
    TELEGRAM_BOT_API_MAX_MEDIA_BYTES,
    TELEGRAM_MEDIA_POLICY_VERSION,
    TELEGRAM_SCRIPTS,
    TYPEFULLY_STATUS,
    X_ASSET_READY,
    X_ASSET_REVIEW_PENDING,
    X_ASSET_STATUS,
    X_PUBLISH_STATUS,
    PublishPipelineError,
    TelegramSendUncertainError,
    _combined_status,
    _details_ready,
    _first_nonempty_line,
    _load_candidates,
    _prepare_telegram_row,
    main,
)
from xpan_auto.copy_review import CopyReviewError
from xpan_auto.video_quality import VideoQualityReport


class PublishPipelineTests(unittest.TestCase):
    def _workbook(self, path: Path, *, typefully_value: str = "") -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "处理结果"
        sheet.append(
            [
                "序号", "处理状态", "标题", "文字介绍", "新夸克链接",
                "Telegram发布状态", "Telegram消息ID", "Telegram发布时间",
                "Telegram错误信息", "Typefully状态",
            ]
        )
        sheet.append(
            [
                1, "SHARE_VALIDATED", "示例课程", "真实资料介绍。",
                "https://pan.quark.cn/s/example", "", "", "", "", typefully_value,
            ]
        )
        workbook.save(path)

    def _video_asset(self, root: Path, *, size: int = 5) -> tuple[Path, dict[str, object]]:
        asset_dir = root / "assets" / "0001"
        asset_dir.mkdir(parents=True)
        original = asset_dir / "project-video-original.mp4"
        with original.open("wb") as handle:
            handle.truncate(size)
        details = {
            "row_index": 1,
            "title": "示例课程",
            "source_share_link": "https://pan.quark.cn/s/example",
            "media_policy_version": TELEGRAM_MEDIA_POLICY_VERSION,
            "publish_media_policy_version": PUBLISH_MEDIA_POLICY_VERSION,
            "media": {
                "kind": "project-video",
                "source_name": "导学.mp4",
                "source_path": str(original),
                "original_path": str(original),
                "size_bytes": size,
            },
            "items": [{"type": "file", "name": "导学.mp4", "extension": ".mp4"}],
        }
        (asset_dir / "telegram-project-details.json").write_text(
            json.dumps(details, ensure_ascii=False), encoding="utf-8"
        )
        row = {
            "序号": 1,
            "标题": "示例课程",
            "文字介绍": "真实资料介绍。",
            "新夸克链接": "https://pan.quark.cn/s/example",
            "Telegram发布状态": "",
            "Telegram消息ID": "",
            "Telegram发布时间": "",
        }
        return original, row

    def _quality(self, path: Path, size: int) -> VideoQualityReport:
        return VideoQualityReport(True, (), str(path), size, 60.0, 1280, 720, "h264", "aac")

    def _captioning_command(self, asset_dir: Path, command) -> str:
        if str(TELEGRAM_SCRIPTS["caption"]) in command:
            (asset_dir / "telegram-caption.html").write_text("Telegram 文案", encoding="utf-8")
            (asset_dir / "x.txt").write_text("X 文案", encoding="utf-8")
        if str(TELEGRAM_SCRIPTS["send"]) in command:
            return '{"mode":"dry-run"}'
        return "{}"

    def _reviewed_files(self, asset_dir: Path) -> None:
        (asset_dir / "telegram-caption.html").write_text("已审阅 Telegram 文案", encoding="utf-8")
        (asset_dir / "x.txt").write_text("已审阅 X 文案", encoding="utf-8")

    def test_first_nonempty_line_selects_first_share_link(self):
        self.assertEqual(
            _first_nonempty_line("\nhttps://pan.quark.cn/s/one\nhttps://pan.quark.cn/s/two"),
            "https://pan.quark.cn/s/one",
        )

    def test_combined_status_requires_real_x_publish_completion(self):
        self.assertEqual(
            _combined_status("已发送", X_ASSET_READY),
            "TELEGRAM_SENT_X_PENDING",
        )
        self.assertEqual(
            _combined_status("已发送", X_ASSET_READY, "已发布"),
            "COMPLETED",
        )
        self.assertEqual(
            _combined_status("已发送", X_ASSET_READY, "", "PUBLISHED"),
            "COMPLETED",
        )
        self.assertEqual(
            _combined_status("已发送", X_ASSET_READY, "待发布", "DRAFT_CREATED"),
            "COMPLETED",
        )
        self.assertEqual(_combined_status("FAILED", X_ASSET_READY), "PARTIAL_FAILED")
        self.assertEqual(
            _combined_status("DRY_RUN", X_ASSET_READY), "TELEGRAM_DRY_RUN_X_READY"
        )
        self.assertEqual(
            _combined_status("等待Web中转", X_ASSET_READY),
            "X_READY_TELEGRAM_PENDING",
        )

    def test_load_candidates_reprocesses_legacy_web_pending_rows(self):
        with TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "pipeline.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(
                [
                    "序号", "处理状态", "Telegram发布状态", "Telegram消息ID",
                    X_ASSET_STATUS, X_PUBLISH_STATUS, TYPEFULLY_STATUS,
                ]
            )
            sheet.append([1, "SHARE_VALIDATED", "已发送", "101", X_ASSET_READY, "已发布", ""])
            sheet.append([2, "SHARE_VALIDATED", "已发送", "102", X_ASSET_READY, "", ""])
            sheet.append([3, "SHARE_VALIDATED", "", "", X_ASSET_READY, "", ""])
            sheet.append([4, "SHARE_VALIDATED", "等待Web中转", "", X_ASSET_READY, "", ""])
            sheet.append([5, "SHARE_VALIDATED", "已发送", "105", X_ASSET_READY, "", "PUBLISHED"])
            sheet.append([6, "SHARE_VALIDATED", "已发送", "106", X_ASSET_READY, "待发布", "DRAFT_CREATED"])
            workbook.save(workbook_path)

            self.assertEqual(
                _load_candidates(
                    workbook_path,
                    only_status={"SHARE_VALIDATED"},
                    limit=10,
                    run_all=True,
                ),
                [3, 4, 5],
            )
            self.assertEqual(
                _load_candidates(
                    workbook_path,
                    only_status={"SHARE_VALIDATED"},
                    limit=10,
                    run_all=True,
                    force_reprocess=True,
                ),
                [2, 3, 4, 5, 6, 7],
            )

    def test_details_ready_requires_only_current_original_video(self):
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            original, _ = self._video_asset(root)
            details = root / "assets" / "0001" / "telegram-project-details.json"
            self.assertTrue(_details_ready(details))
            self.assertEqual(list(original.parent.glob("project-video*.mp4")), [original])

    def test_details_ready_rejects_old_three_video_policy(self):
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            original, _ = self._video_asset(root)
            details = original.parent / "telegram-project-details.json"
            payload = json.loads(details.read_text(encoding="utf-8"))
            payload["publish_media_policy_version"] = PUBLISH_MEDIA_POLICY_VERSION - 1
            details.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            self.assertFalse(_details_ready(details))

    def test_details_ready_rejects_truncated_project_tree(self):
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            original, _ = self._video_asset(root)
            details = original.parent / "telegram-project-details.json"
            payload = json.loads(details.read_text(encoding="utf-8"))
            payload["summary"] = {"truncated": True}
            details.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

            self.assertFalse(_details_ready(details))

    def test_small_video_uses_original_bot_api_and_prepares_x(self):
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            original, row = self._video_asset(root)
            asset_dir = original.parent
            report = self._quality(original, original.stat().st_size)
            with patch(
                "xpan_auto.publish_pipeline.inspect_publish_video", return_value=report
            ), patch(
                "xpan_auto.publish_pipeline._run_command",
                side_effect=lambda command: self._captioning_command(asset_dir, command),
            ) as run_command:
                result = _prepare_telegram_row(
                    root / "pipeline.xlsx", 2, row,
                    image_root=root / "assets",
                    quark_db_path=root / "quark.sqlite3",
                    cookie_file=root / "cookie.json",
                    send=False, force_send=False, image_column=1,
                )

            commands = [list(call.args[0]) for call in run_command.call_args_list]
            send_command = next(c for c in commands if str(TELEGRAM_SCRIPTS["send"]) in c)
            caption_command = next(c for c in commands if str(TELEGRAM_SCRIPTS["caption"]) in c)
            self.assertEqual(Path(result["media"]), original)
            self.assertEqual(Path(result["x_media"]), original)
            self.assertEqual(result["x_status"], X_ASSET_REVIEW_PENDING)
            self.assertEqual(result["telegram_route"], "bot-api")
            self.assertIn("--video", send_command)
            self.assertNotIn("--transport", send_command)
            self.assertIn("--x-output", caption_command)

    def test_exactly_50mb_final_video_is_blocked_before_upload(self):
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            workbook_path = root / "pipeline.xlsx"
            self._workbook(workbook_path)
            original, row = self._video_asset(root, size=TELEGRAM_BOT_API_MAX_MEDIA_BYTES)
            asset_dir = original.parent
            self._reviewed_files(asset_dir)
            report = self._quality(original, original.stat().st_size)
            with patch(
                "xpan_auto.publish_pipeline.inspect_publish_video", return_value=report
            ), patch(
                "xpan_auto.publish_pipeline._run_command",
                side_effect=lambda command: self._captioning_command(asset_dir, command),
            ) as run_command:
                with self.assertRaisesRegex(PublishPipelineError, "严格小于 50,000,000 bytes"):
                    _prepare_telegram_row(
                        workbook_path, 2, row,
                        image_root=root / "assets",
                        quark_db_path=root / "quark.sqlite3",
                        cookie_file=root / "cookie.json",
                        send=True, force_send=False, image_column=1,
                    )

            commands = [list(call.args[0]) for call in run_command.call_args_list]
            self.assertFalse(any(str(TELEGRAM_SCRIPTS["send"]) in c for c in commands))

    def test_bot_api_uncertainty_keeps_sending_status(self):
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            workbook_path = root / "pipeline.xlsx"
            self._workbook(workbook_path)
            original, row = self._video_asset(root)
            asset_dir = original.parent
            self._reviewed_files(asset_dir)
            report = self._quality(original, original.stat().st_size)

            def command_result(command):
                self._captioning_command(asset_dir, command)
                if str(TELEGRAM_SCRIPTS["send"]) in command:
                    raise PublishPipelineError("upload process failed")
                return "{}"

            with patch(
                "xpan_auto.publish_pipeline.inspect_publish_video", return_value=report
            ), patch("xpan_auto.publish_pipeline._run_command", side_effect=command_result), patch(
                "xpan_auto.publish_pipeline.validate_review", return_value={"style_label": "疑问／观点型"}
            ):
                with self.assertRaises(TelegramSendUncertainError):
                    _prepare_telegram_row(
                        workbook_path, 2, row,
                        image_root=root / "assets",
                        quark_db_path=root / "quark.sqlite3",
                        cookie_file=root / "cookie.json",
                        send=True, force_send=False, image_column=1,
                    )

            saved = load_workbook(workbook_path)
            headers = {cell.value: cell.column for cell in saved.active[1]}
            self.assertEqual(saved.active.cell(2, headers["Telegram发布状态"]).value, "发送中")
            self.assertEqual(saved.active.cell(2, headers[X_ASSET_STATUS]).value, X_ASSET_READY)
            saved.close()

    def test_main_writes_x_ready_and_preserves_typefully_history(self):
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            workbook_path = root / "pipeline.xlsx"
            self._workbook(workbook_path, typefully_value="DRAFT_CREATED_OLD")
            original = root / "assets" / "0001" / "project-video-original.mp4"
            original.parent.mkdir(parents=True)
            original.write_bytes(b"video")
            result = {
                "status": "DRY_RUN", "message_id": "", "published_at": "",
                "media_type": "video", "media": str(original),
                "original_media": str(original), "x_media": str(original),
                "x_status": X_ASSET_READY, "telegram_route": "bot-api",
                "caption": str(original.parent / "telegram-caption.html"),
            }
            with patch(
                "xpan_auto.publish_pipeline.FIXED_WORKBOOK_PATH", workbook_path.resolve()
            ), patch(
                "xpan_auto.publish_pipeline._prepare_telegram_row", return_value=result
            ):
                self.assertEqual(
                    main(["--xlsx", str(workbook_path), "--asset-root", str(root / "assets")]),
                    0,
                )

            saved = load_workbook(workbook_path)
            headers = {cell.value: cell.column for cell in saved.active[1]}
            self.assertEqual(saved.active.cell(2, headers[X_ASSET_STATUS]).value, X_ASSET_READY)
            self.assertEqual(saved.active.cell(2, headers["Typefully状态"]).value, "DRAFT_CREATED_OLD")
            self.assertEqual(saved.active.cell(2, headers["原视频"]).value, str(original))
            self.assertEqual(saved.active.cell(2, headers["Telegram视频"]).value, str(original))
            saved.close()

    def test_main_records_large_image_failure_without_web_relay_exit(self):
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            workbook_path = root / "pipeline.xlsx"
            self._workbook(workbook_path)
            with patch(
                "xpan_auto.publish_pipeline.FIXED_WORKBOOK_PATH", workbook_path.resolve()
            ), patch(
                "xpan_auto.publish_pipeline._run_command", return_value="{}"
            ), patch(
                "xpan_auto.publish_pipeline._prepare_telegram_row",
                side_effect=PublishPipelineError(
                    "最终图片必须严格小于 50,000,000 bytes；Telegram 只允许 Bot API"
                ),
            ), patch(
                "xpan_auto.publish_pipeline.validate_review",
                return_value={"style_label": "疑问／观点型", "opening": "不重复的开头", "layout": "blocks-2"},
            ):
                self.assertEqual(
                    main(
                        [
                            "--xlsx", str(workbook_path),
                            "--asset-root", str(root / "assets"),
                            "--telegram-send",
                        ]
                    ),
                    0,
                )

            saved = load_workbook(workbook_path)
            headers = {cell.value: cell.column for cell in saved.active[1]}
            self.assertEqual(saved.active.cell(2, headers["Telegram发布状态"]).value, "FAILED")
            self.assertIn(
                "Telegram 只允许 Bot API",
                saved.active.cell(2, headers["Telegram错误信息"]).value,
            )
            saved.close()

    def test_real_send_preflight_uses_bot_api_check_before_rows(self):
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            workbook_path = root / "pipeline.xlsx"
            self._workbook(workbook_path)
            with patch(
                "xpan_auto.publish_pipeline.FIXED_WORKBOOK_PATH", workbook_path.resolve()
            ), patch(
                "xpan_auto.publish_pipeline._run_command",
                side_effect=PublishPipelineError("Bot API check failed"),
            ) as run_command, patch(
                "xpan_auto.publish_pipeline._prepare_telegram_row"
            ) as prepare_row:
                with self.assertRaisesRegex(PublishPipelineError, "Bot API check failed"):
                    main(
                        [
                            "--xlsx", str(workbook_path),
                            "--asset-root", str(root / "assets"),
                            "--telegram-send",
                        ]
                    )
            prepare_row.assert_not_called()
            command = list(run_command.call_args.args[0])
            self.assertEqual(command[-1], "--check")
            self.assertNotIn("mtproto", " ".join(command).lower())

    def test_real_send_requires_editorial_review_before_preparing_any_row(self):
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            workbook_path = root / "pipeline.xlsx"
            self._workbook(workbook_path)
            with patch(
                "xpan_auto.publish_pipeline.FIXED_WORKBOOK_PATH", workbook_path.resolve()
            ), patch(
                "xpan_auto.publish_pipeline._run_command", return_value="{}"
            ) as run_command, patch(
                "xpan_auto.publish_pipeline._prepare_telegram_row"
            ) as prepare_row:
                with self.assertRaises(CopyReviewError):
                    main(["--xlsx", str(workbook_path), "--telegram-send", "--row-number", "2"])
            prepare_row.assert_not_called()
            self.assertEqual(len(run_command.call_args_list), 1)
            saved = load_workbook(workbook_path)
            self.assertEqual(saved.active.cell(2, 6).value, None)
            saved.close()

    def test_batch_review_failure_blocks_first_telegram_send(self):
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            workbook_path = root / "pipeline.xlsx"
            self._workbook(workbook_path)
            workbook = load_workbook(workbook_path)
            workbook.active.append([
                2, "SHARE_VALIDATED", "第二条教程", "第二条真实介绍。",
                "https://pan.quark.cn/s/second", "", "", "", "", "",
            ])
            workbook.save(workbook_path)
            workbook.close()
            with patch(
                "xpan_auto.publish_pipeline.FIXED_WORKBOOK_PATH", workbook_path.resolve()
            ), patch(
                "xpan_auto.publish_pipeline._run_command", return_value="{}"
            ), patch(
                "xpan_auto.publish_pipeline.validate_review",
                side_effect=[
                    {"style_label": "指定人群型", "opening": "独立开头", "layout": "blocks-1"},
                    CopyReviewError("第二条尚未审阅"),
                ],
            ) as validate, patch(
                "xpan_auto.publish_pipeline._prepare_telegram_row"
            ) as prepare_row:
                with self.assertRaisesRegex(CopyReviewError, "第二条尚未审阅"):
                    main([
                        "--xlsx", str(workbook_path), "--telegram-send",
                        "--row-number", "2", "--row-number", "3",
                    ])
            self.assertEqual(validate.call_count, 2)
            prepare_row.assert_not_called()

    def test_real_send_reuses_reviewed_copy_without_regenerating(self):
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            workbook_path = root / "pipeline.xlsx"
            self._workbook(workbook_path)
            original, row = self._video_asset(root)
            asset_dir = original.parent
            self._reviewed_files(asset_dir)
            report = self._quality(original, original.stat().st_size)
            with patch(
                "xpan_auto.publish_pipeline.inspect_publish_video", return_value=report
            ), patch(
                "xpan_auto.publish_pipeline.validate_review",
                return_value={"style_label": "疑问／观点型", "opening": "独立开头", "layout": "blocks-2"},
            ), patch(
                "xpan_auto.publish_pipeline._run_command", return_value='{"message_id": 123}'
            ) as run_command:
                result = _prepare_telegram_row(
                    workbook_path, 2, row,
                    image_root=root / "assets", quark_db_path=root / "quark.sqlite3",
                    cookie_file=root / "cookie.json", send=True, force_send=False, image_column=1,
                )
            commands = [list(call.args[0]) for call in run_command.call_args_list]
            self.assertEqual(result["message_id"], "123")
            self.assertFalse(any(str(TELEGRAM_SCRIPTS["caption"]) in c for c in commands))
            self.assertEqual((asset_dir / "x.txt").read_text(encoding="utf-8"), "已审阅 X 文案")

    def test_dry_retry_preserves_existing_reviewed_copy(self):
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            workbook_path = root / "pipeline.xlsx"
            self._workbook(workbook_path)
            original, row = self._video_asset(root)
            asset_dir = original.parent
            self._reviewed_files(asset_dir)
            (asset_dir / "copy-review.json").write_text("{}", encoding="utf-8")
            row["Telegram发布状态"] = "已发送"
            row["Telegram消息ID"] = "123"
            report = self._quality(original, original.stat().st_size)
            with patch(
                "xpan_auto.publish_pipeline.inspect_publish_video", return_value=report
            ), patch(
                "xpan_auto.publish_pipeline._run_command"
            ) as run_command:
                result = _prepare_telegram_row(
                    workbook_path, 2, row,
                    image_root=root / "assets", quark_db_path=root / "quark.sqlite3",
                    cookie_file=root / "cookie.json", send=False, force_send=False,
                    image_column=1,
                )
            run_command.assert_not_called()
            self.assertEqual(result["message_id"], "123")
            self.assertEqual((asset_dir / "x.txt").read_text(encoding="utf-8"), "已审阅 X 文案")

    def test_parser_default_is_the_fixed_workflow_workbook(self):
        from xpan_auto.publish_pipeline import _parser

        self.assertEqual(Path(_parser().parse_args([]).xlsx), FIXED_WORKBOOK_PATH)

    def test_main_rejects_alternate_workbook_before_reading_or_sending(self):
        with TemporaryDirectory() as temp_dir:
            alternate = Path(temp_dir) / "alternate.xlsx"
            self._workbook(alternate)
            with patch("xpan_auto.publish_pipeline._load_candidates") as load_candidates, patch(
                "xpan_auto.publish_pipeline._run_command"
            ) as run_command:
                with self.assertRaisesRegex(PublishPipelineError, "fixed at"):
                    main(["--xlsx", str(alternate), "--telegram-send"])
            load_candidates.assert_not_called()
            run_command.assert_not_called()


if __name__ == "__main__":
    unittest.main()
