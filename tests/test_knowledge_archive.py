import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from xpan_auto.excel_exporter import export_rows
from xpan_auto.knowledge_flow import (
    DEFAULT_LEDGER_XLSX,
    Entry,
    KnowledgeState,
    LEDGER_COLUMNS,
    _ledger_sequence_by_share_id,
    _next_archive_index,
    _run_quark_for_entry,
    ensure_ledger,
    upsert_ledger_row,
)
from xpan_auto.quark_adapter import NewShare, QuarkItem, SaveResult, UsageSummary


class FakeQuarkAdapter:
    def inspect_share(self, _url: str) -> None:
        return None

    def ensure_folder(self, _path: str) -> str:
        return "folder-1"

    def save_share(self, _url: str, _folder_id: str) -> SaveResult:
        return SaveResult(
            task_id="task-1",
            target_folder_id="folder-1",
            items=(QuarkItem(fid="file-1", name="course.zip", is_dir=False, size_bytes=1536),),
        )

    def list_folder(self, _folder_id: str):
        return []

    def calculate_usage(self, _items) -> UsageSummary:
        return UsageSummary(total_bytes=1536, file_count=1, folder_count=0)

    def create_share(self, item: QuarkItem, **_kwargs) -> NewShare:
        return NewShare(fid=item.fid, title=item.name, share_url="https://pan.quark.cn/s/new-link")

    def validate_share(self, _share: NewShare) -> bool:
        return True


class KnowledgeArchiveTests(unittest.TestCase):
    def test_default_ledger_is_existing_publish_workbook(self) -> None:
        self.assertEqual(DEFAULT_LEDGER_XLSX, "output/quark_links.typefully.xlsx")

    def test_ledger_validation_preserves_existing_extra_columns(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            xlsx = Path(directory) / "ledger.xlsx"
            export_rows(
                [{"序号": 1, "原夸克链接": "https://pan.quark.cn/s/original-a"}],
                xlsx,
            )
            workbook = load_workbook(xlsx)
            workbook.active.cell(row=1, column=17, value="Telegram发布状态")
            workbook.active.cell(row=2, column=17, value="已发送")
            workbook.save(xlsx)
            workbook.close()

            ensure_ledger(xlsx)
            workbook = load_workbook(xlsx)
            try:
                headers = [cell.value for cell in workbook.active[1]]
            finally:
                workbook.close()

        self.assertEqual(headers[: len(LEDGER_COLUMNS)], LEDGER_COLUMNS)
        self.assertEqual(headers[16], "Telegram发布状态")

    def test_new_row_appends_and_copies_previous_row_format(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            xlsx = Path(directory) / "ledger.xlsx"
            export_rows(
                [{"序号": 58, "原夸克链接": "https://pan.quark.cn/s/original-a"}],
                xlsx,
            )
            workbook = load_workbook(xlsx)
            sheet = workbook.active
            sheet.cell(row=1, column=17, value="Telegram发布状态")
            sheet.cell(row=2, column=17, value="已发送")
            sheet["A2"].fill = PatternFill("solid", fgColor="FFF2CC")
            sheet.row_dimensions[2].height = 33
            workbook.save(xlsx)
            workbook.close()

            row_number = upsert_ledger_row(
                xlsx,
                {
                    "序号": 59,
                    "资料标题": "资料 B",
                    "原夸克链接": "https://pan.quark.cn/s/original-b",
                    "新夸克链接": "https://pan.quark.cn/s/new-b",
                    "保存目录": "/推文分享/0059_资料 B",
                    "文件数量": 1,
                    "总大小": "1.50 KB",
                    "处理状态": "SHARE_VALIDATED",
                },
            )

            workbook = load_workbook(xlsx)
            try:
                sheet = workbook.active
                self.assertEqual(row_number, 3)
                self.assertEqual(sheet["A3"].value, 59)
                self.assertEqual(sheet["C3"].value, "资料 B")
                self.assertEqual(sheet["K3"].value, "1.50 KB")
                self.assertEqual(sheet.cell(row=2, column=17).value, "已发送")
                self.assertIsNone(sheet.cell(row=3, column=17).value)
                self.assertEqual(sheet["A3"].style_id, sheet["A2"].style_id)
                self.assertEqual(sheet.row_dimensions[3].height, 33)
            finally:
                workbook.close()

    def test_next_folder_index_uses_state_ledger_and_cloud_root(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            state = KnowledgeState(root / "state.sqlite3")
            entry = Entry("08/20", "资料 A", "https://pan.quark.cn/s/original-a")
            state.upsert_archived(
                entry,
                category="工具分享",
                save_path="/推文分享/0024_资料 A",
                status="SHARE_VALIDATED",
                updated_at="2026-08-20T20:00:00+08:00",
            )
            xlsx = root / "ledger.xlsx"
            export_rows(
                [
                    {
                        "序号": 25,
                        "资料标题": "资料 B",
                        "原夸克链接": "https://pan.quark.cn/s/original-b",
                        "保存目录": "/推文分享/0025_资料 B",
                    },
                    {
                        "序号": 31,
                        "资料标题": "失败待重试",
                        "原夸克链接": "https://pan.quark.cn/s/original-failed",
                        "处理状态": "FAILED",
                    },
                ],
                xlsx,
            )
            root_items = [
                SimpleNamespace(name="0026_资料 C", is_dir=True),
                SimpleNamespace(name="未编号目录", is_dir=True),
            ]

            next_index = _next_archive_index(state, xlsx, root_items)

        self.assertEqual(next_index, 32)

    def test_retry_uses_existing_ledger_sequence(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            xlsx = Path(directory) / "ledger.xlsx"
            export_rows(
                [{
                    "序号": 60,
                    "原夸克链接": "https://pan.quark.cn/s/original-failed",
                    "处理状态": "FAILED",
                }],
                xlsx,
            )

            sequences = _ledger_sequence_by_share_id(xlsx)

        self.assertEqual(sequences["original-failed"], 60)

    def test_quark_retry_reuses_existing_saved_items(self) -> None:
        class ExistingFolderAdapter(FakeQuarkAdapter):
            def list_folder(self, _folder_id: str):
                return [QuarkItem(fid="file-existing", name="existing.zip", is_dir=False, size_bytes=2048)]

            def save_share(self, _url: str, _folder_id: str) -> SaveResult:
                raise AssertionError("existing retry folder must not be saved again")

            def calculate_usage(self, _items) -> UsageSummary:
                return UsageSummary(total_bytes=2048, file_count=1, folder_count=0)

        args = SimpleNamespace(
            target_root="/推文分享",
            share_expiry="permanent",
            share_password_mode="none",
            share_password=None,
        )

        result = _run_quark_for_entry(
            ExistingFolderAdapter(),
            Entry("08/20", "课程", "https://pan.quark.cn/s/original-a"),
            60,
            args,
        )

        self.assertEqual(result[0], "SHARE_VALIDATED")
        self.assertEqual(result[3], 1)
        self.assertEqual(result[4], "2.00 KB")

    def test_quark_run_returns_total_saved_size(self) -> None:
        args = SimpleNamespace(
            target_root="/推文分享",
            share_expiry="permanent",
            share_password_mode="none",
            share_password=None,
        )

        status, new_url, passcode, file_count, size, error = _run_quark_for_entry(
            FakeQuarkAdapter(),
            Entry("08/20", "课程", "https://pan.quark.cn/s/original-a"),
            27,
            args,
        )

        self.assertEqual(status, "SHARE_VALIDATED")
        self.assertEqual(new_url, "https://pan.quark.cn/s/new-link")
        self.assertEqual(passcode, "")
        self.assertEqual(file_count, 1)
        self.assertEqual(size, "1.50 KB")
        self.assertEqual(error, "")


if __name__ == "__main__":
    unittest.main()
