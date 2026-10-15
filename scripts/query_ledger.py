#!/usr/bin/env python3
"""只读查询固定发布主表，按 Excel 表头输出 JSON。"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = PROJECT_ROOT / "output" / "quark_links.typefully.xlsx"
SHEET_NAME = "处理结果"
SEQUENCE_FIELD = "序号"


def _parse_sequence_spec(value: str) -> set[int]:
    """Parse comma-separated sequence numbers and inclusive ranges."""
    result: set[int] = set()
    for part in value.split(","):
        part = part.strip()
        if not part:
            raise ValueError("序号筛选中存在空项")
        if "-" in part:
            pieces = part.split("-")
            if len(pieces) != 2 or not all(piece.strip().isdigit() for piece in pieces):
                raise ValueError(f"无效序号范围：{part!r}（示例：1-5,8）")
            start, end = map(lambda item: int(item.strip()), pieces)
            if start < 1 or end < start:
                raise ValueError(f"无效序号范围：{part!r}；序号须为正数且起点不大于终点")
            result.update(range(start, end + 1))
        else:
            if not part.isdigit() or int(part) < 1:
                raise ValueError(f"无效序号：{part!r}；序号须为正整数")
            result.add(int(part))
    return result


def _parse_where(items: list[str], headers: list[str]) -> list[tuple[str, str]]:
    filters = []
    for item in items:
        if "=" not in item:
            raise ValueError(f"精确筛选格式应为 字段=值：{item!r}")
        field, value = item.split("=", 1)
        field = field.strip()
        if not field:
            raise ValueError(f"精确筛选缺少字段名：{item!r}")
        if field not in headers:
            raise ValueError(f"未知字段：{field}；可用字段：{', '.join(headers)}")
        filters.append((field, value.strip()))
    return filters


def query_ledger(args: argparse.Namespace) -> dict[str, Any]:
    if not WORKBOOK_PATH.is_file():
        raise FileNotFoundError(f"主表不存在：{WORKBOOK_PATH}")
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"主表缺少工作表：{SHEET_NAME}")
        sheet = workbook[SHEET_NAME]
        header_values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [str(value).strip() if value is not None else "" for value in header_values]
        header_to_index: dict[str, int] = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            if header in header_to_index:
                raise ValueError(f"表头重复：{header}")
            header_to_index[header] = index
        if SEQUENCE_FIELD not in header_to_index:
            raise ValueError(f"主表缺少必要字段：{SEQUENCE_FIELD}")

        selected_columns = None
        if args.columns:
            selected_columns = [name.strip() for name in args.columns.split(",") if name.strip()]
            missing = [name for name in selected_columns if name not in header_to_index]
            if missing:
                raise ValueError(f"未知列：{', '.join(missing)}；可用字段：{', '.join(header_to_index)}")
            if not selected_columns:
                raise ValueError("列选择不能为空")
        where = _parse_where(args.where, list(header_to_index))
        sequence_set = _parse_sequence_spec(args.sequence) if args.sequence else None

        data_rows = []
        nonempty_rows = 0
        for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            record = {name: values[index] if index < len(values) else None
                      for name, index in header_to_index.items()}
            if all(value is None or value == "" for value in record.values()):
                continue
            nonempty_rows += 1
            sequence = record.get(SEQUENCE_FIELD)
            if sequence_set is not None:
                try:
                    if int(sequence) not in sequence_set:
                        continue
                except (TypeError, ValueError):
                    continue
            if args.row_number is not None and excel_row != args.row_number:
                continue
            if args.contains and not any(args.contains.casefold() in str(value).casefold()
                                         for value in record.values() if value is not None):
                continue
            if any(str(record[field] if record[field] is not None else "") != value
                   for field, value in where):
                continue
            fields = selected_columns or list(header_to_index)
            output = {field: record[field] for field in fields}
            output["序号"] = sequence
            data_rows.append({"excel_row": excel_row, "sequence": sequence, "data": output})

        matched_count = len(data_rows)
        summary_only = args.summary or not any((args.sequence, args.row_number, args.columns,
                                                 args.contains, args.where))
        limit = args.limit if args.limit is not None else 20
        returned_rows = 0 if summary_only else min(matched_count, limit)
        truncated = False if summary_only else matched_count > returned_rows
        data_rows = [] if summary_only else data_rows[:limit]
        try:
            workbook_label = str(WORKBOOK_PATH.relative_to(PROJECT_ROOT))
        except ValueError:  # enables isolated tests without changing query behavior
            workbook_label = WORKBOOK_PATH.name
        summary = {
            "workbook": workbook_label,
            "sheet": SHEET_NAME,
            "headers": list(header_to_index),
            "data_rows": nonempty_rows,
            "matched_rows": matched_count,
            "returned_rows": returned_rows,
            "truncated": truncated,
        }
        return {"summary": summary, "rows": data_rows}
    finally:
        workbook.close()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="按表头只读查询固定发布主表，输出 JSON。")
    parser.add_argument("--sequence", "--seq", help="序号，例如 3、2-5 或 2-5,8")
    parser.add_argument("--row-number", "--row", type=int, help="真实 Excel 行号（表头为第 1 行）")
    parser.add_argument("--columns", help="只输出指定字段，逗号分隔")
    parser.add_argument("--contains", help="在整行字段中不区分大小写的关键词查找")
    parser.add_argument("--where", action="append", default=[], metavar="字段=值",
                        help="字段精确匹配；可重复指定")
    parser.add_argument("--limit", type=int, help="最多返回条数")
    parser.add_argument("--summary", action="store_true", help="只返回摘要，不返回记录")
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    if args.row_number is not None and args.row_number < 2:
        parser.error("--row-number 必须是 2 或更大的 Excel 数据行号")
    if args.limit is not None and args.limit < 1:
        parser.error("--limit 必须是正整数")
    try:
        result = query_ledger(args)
    except (FileNotFoundError, ValueError) as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 2
    print(json.dumps(result, ensure_ascii=False, default=str, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
