#!/usr/bin/env python3
"""Rebase existing asset paths; back up JSON and plan workbook-only edits."""
from __future__ import annotations
import argparse
import json
import re
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from xpan_auto.runtime import resolve_project_path

PATH_COLUMNS = {"图片", "视频", "原视频", "Telegram视频", "媒体路径", "Typefully媒体路径", "X文案路径", "Telegram文案路径"}


def replacement(value):
    if not isinstance(value, str) or not (re.match(r"^[A-Za-z]:[\\/]", value) or "\\" in value or value.startswith("/")):
        return value, False
    resolved = resolve_project_path(value)
    if not resolved.is_relative_to(ROOT) or resolved == Path(value):
        return value, False
    if resolved.is_file():
        return str(resolved), False
    return value, True


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--execute", action="store_true")
    parser.add_argument("--report-dir", type=Path, required=True)
    args = parser.parse_args()
    args.report_dir.mkdir(parents=True, exist_ok=True)
    changes, missing = [], []

    def walk(value, filename, key=""):
        if isinstance(value, dict):
            return {k: walk(v, filename, key + "." + k) for k, v in value.items()}
        if isinstance(value, list):
            return [walk(v, filename, key + f"[{i}]") for i, v in enumerate(value)]
        new, unresolved = replacement(value)
        if unresolved:
            missing.append({"file": filename, "key": key, "path": value})
        if new != value:
            changes.append({"file": filename, "key": key, "before": value, "after": new})
        return new

    json_files = 0
    for file in sorted((ROOT / "output/assets").glob("*/*.json")):
        original = json.loads(file.read_text(encoding="utf-8-sig"))
        updated = walk(original, str(file.relative_to(ROOT)))
        if updated == original:
            continue
        json_files += 1
        if args.execute:
            backup = args.report_dir / "metadata-before" / file.relative_to(ROOT)
            if backup.exists():
                raise RuntimeError(f"Backup already exists; refusing to overwrite: {backup}")
            backup.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(file, backup)
            temporary = file.with_suffix(".json.migrating")
            temporary.write_text(json.dumps(updated, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            temporary.replace(file)

    # The workbook is only read here. Apply the bounded plan with the spreadsheet
    # authoring workflow, preserving all non-path cells and native features.
    from openpyxl import load_workbook
    workbook_path = ROOT / "output/quark_links.typefully.xlsx"
    workbook = load_workbook(workbook_path)
    sheet = workbook["处理结果"]
    edits = []
    for column in sheet[1]:
        if column.value not in PATH_COLUMNS:
            continue
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row, column.column)
            new, unresolved = replacement(cell.value)
            if unresolved:
                missing.append({"file": str(workbook_path.relative_to(ROOT)), "key": cell.coordinate, "path": cell.value})
            if new != cell.value:
                edits.append({"cell": cell.coordinate, "before": cell.value, "after": new})
    workbook.close()
    report = {"executed": args.execute, "json_files": json_files, "json_changes": changes,
              "workbook": str(workbook_path), "sheet": "处理结果", "workbook_edits": edits,
              "unresolved": missing}
    report_name = "path-migration-applied.json" if args.execute else "path-migration.json"
    (args.report_dir / report_name).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"executed": args.execute, "json_files": json_files, "json_fields": len(changes),
                      "workbook_edits_planned": len(edits), "unresolved_references": len(missing)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
